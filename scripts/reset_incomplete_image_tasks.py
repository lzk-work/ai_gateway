"""Reset only non-success image task state for incomplete SKUs in one batch."""

from __future__ import annotations

import argparse
import json
import os
from collections import Counter
from pathlib import Path
from typing import Any

from ai_gateway.subtasks.mxapi_generate_images import is_completed_success


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def key(row: dict[str, Any]) -> str:
    return f"{row.get('sku')}::{row.get('image_name')}"


def latest_by_key(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {key(row): row for row in rows if row.get("sku") and row.get("image_name")}


def write_jsonl_atomic(path: Path, rows: list[dict[str, Any]]) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8", newline="\n") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(temporary, path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("batch_root", type=Path)
    parser.add_argument("--desired-count", type=int, default=5)
    parser.add_argument("--expected-incomplete", type=int)
    parser.add_argument("--audit-path", type=Path, required=True)
    args = parser.parse_args()

    image_dir = args.batch_root / "04_generate_images"
    checkpoint_path = image_dir / "image_generation_checkpoint.jsonl"
    results_path = image_dir / "image_generation_results.jsonl"

    checkpoint_latest = latest_by_key(read_jsonl(checkpoint_path))
    result_latest = latest_by_key(read_jsonl(results_path))
    combined = dict(checkpoint_latest)
    combined.update(result_latest)

    success_by_sku: Counter[str] = Counter()
    success_keys: set[str] = set()
    for row_key, row in combined.items():
        if is_completed_success(row):
            sku = str(row["sku"]).strip()
            success_by_sku[sku] += 1
            success_keys.add(row_key)

    # The image input is authoritative for which SKUs belong to the stage.
    from openpyxl import load_workbook

    input_path = args.batch_root / "03_build_image_input" / "walmart_sub_image_input_result.xlsx"
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    sheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.active
    headers = {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value}
    sku_header = headers["SKU"]
    all_skus: set[str] = set()
    for values in sheet.iter_rows(
        min_row=2,
        min_col=sku_header,
        max_col=sku_header,
        values_only=True,
    ):
        if values[0] is not None and str(values[0]).strip():
            all_skus.add(str(values[0]).strip())
    incomplete = sorted(sku for sku in all_skus if success_by_sku[sku] < args.desired_count)
    if args.expected_incomplete is not None and len(incomplete) != args.expected_incomplete:
        raise RuntimeError(
            f"Refusing reset: expected {args.expected_incomplete} incomplete SKUs, found {len(incomplete)}"
        )

    def retained(rows_by_key: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
        return [
            row
            for row_key, row in rows_by_key.items()
            if str(row.get("sku") or "").strip() not in incomplete or row_key in success_keys
        ]

    checkpoint_out = retained(checkpoint_latest)
    results_out = retained(result_latest)
    removed_checkpoint = len(checkpoint_latest) - len(checkpoint_out)
    removed_results = len(result_latest) - len(results_out)
    pending = sum(args.desired_count - success_by_sku[sku] for sku in incomplete)

    audit = {
        "batch_root": str(args.batch_root.resolve()),
        "desired_count": args.desired_count,
        "incomplete_sku_count": len(incomplete),
        "incomplete_skus": incomplete,
        "preserved_success_count": len(success_keys),
        "pending_generation_count": pending,
        "removed_non_success_checkpoint_records": removed_checkpoint,
        "removed_non_success_result_records": removed_results,
    }
    args.audit_path.parent.mkdir(parents=True, exist_ok=True)
    args.audit_path.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    write_jsonl_atomic(checkpoint_path, checkpoint_out)
    write_jsonl_atomic(results_path, results_out)
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
