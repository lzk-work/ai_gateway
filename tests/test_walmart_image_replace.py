"""Offline replacement acceptance; all remote calls are forbidden or mocked."""
import base64
import importlib.util
import json
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock
import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
spec = importlib.util.spec_from_file_location('replace_workflow_test', ROOT / 'subtasks/walmart_image_replace/replace_workflow_common.py')
w = importlib.util.module_from_spec(spec)
spec.loader.exec_module(w)
OSS = 'https://test-bucket.oss-cn-beijing.aliyuncs.com/images/walmart'
REF = 'https://i5.walmartimages.com/asr/reference.jpeg'
PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=')

def row(main=True, subs=1, source='SOURCE', result='RESULT', old=0):
    value = {'来源SKU': source, '结果SKU': result, '店铺': 'store', '标题': 'Bottle', '五点': 'Steel\nPortable',
        '参考主图链接': REF, '参考副图链接1': REF.replace('reference','secondary'),
        '已有主图链接': f'{OSS}/EARLIEST/new_main2_EARLIEST.png' if main else f'{OSS}/{source}/main_{source}.jpg'}
    for i in range(1, subs + 1):
        value[f'已有副图链接{i}'] = f'{OSS}/{source}/new_sub{i}_{source}.png'
    for i in range(subs + 1, subs + old + 1):
        value[f'已有副图链接{i}'] = f'{OSS}/{source}/sub{i}_{source}.jpg'
    return value

@pytest.fixture
def workspace(tmp_path, monkeypatch):
    source = tmp_path / 'input.xlsx'
    cfg = {'image_provider':'tuzi', 'input':{'excel_path':str(source),'sheet_name':'Sheet1'},
        'execution':{'max_records':100,'concurrency':2,'image_concurrency':2,'oss_concurrency':2},
        'workflow':{}}
    monkeypatch.setattr(w,'BATCHES_ROOT',tmp_path / 'batches')
    monkeypatch.setattr(w,'load_task_config',lambda:cfg)
    monkeypatch.setattr(w,'oss_settings',lambda:{'bucket':'test-bucket','endpoint':'oss-cn-beijing.aliyuncs.com','default_prefix':'images'})
    import requests
    monkeypatch.setattr(requests.sessions.Session,'request',Mock(side_effect=AssertionError('network forbidden')))
    import urllib.request
    monkeypatch.setattr(urllib.request,'urlopen',Mock(side_effect=AssertionError('network forbidden')))
    def write(rows):
        wb = Workbook()
        wb.active.title = 'Sheet1'
        headers = list(dict.fromkeys(key for r in rows for key in r))
        wb.active.append(headers)
        for r in rows:
            wb.active.append([r.get(h,'') for h in headers])
        wb.save(source)
    return cfg, write

@pytest.mark.parametrize('main,subs,expected',[(True,1,(0,3)),(False,0,(0,4)),(False,5,(0,0)),(True,6,(0,0)),(True,3,(0,1)),(True,4,(0,0))])
def test_dynamic_counts_and_history_reuse(workspace, main, subs, expected):
    cfg, write = workspace
    write([row(main,subs)])
    r = w.prepare()[0]
    assert not r['errors']
    assert (r['need_main'],r['need_sub']) == expected
    assert r['oss_directory']['prefix'] == 'images/walmart/SOURCE'
    assert len(r['tasks']) == sum(expected)
    assert all('_SOURCE_' in t['image_name'] for t in r['tasks'])
    assert all('RESULT' not in t['image_name'] for t in r['tasks'])


def test_source_directory_when_all_urls_use_other_history_skus(workspace):
    _, write = workspace
    r = row(False,2)
    r['已有主图链接'] = f'{OSS}/FIRST/main_FIRST.jpg'
    r['已有副图链接1'] = f'{OSS}/SECOND/new_sub1_SECOND.png'
    r['已有副图链接2'] = f'{OSS}/SECOND/new_sub2_SECOND.png'
    write([r])
    planned = w.prepare()[0]
    assert not planned['errors']
    assert planned['oss_directory']['prefix'] == 'images/walmart/SOURCE'
    assert len(planned['tasks']) == 2


def test_stable_names_and_same_source_independent_products(workspace):
    _, write = workspace
    write([row(result='B'),row(result='C')])
    first = w.prepare()
    second = w.prepare()
    assert first == second
    assert first[0]['record_id'] != first[1]['record_id']
    names = [t['image_name'] for r in first for t in r['tasks']]
    assert len(names) == len(set(names)) == 6


def test_duplicate_result_rows_block_both(workspace):
    _, write = workspace
    write([row(),row()])
    assert all(r['errors'] and not r['tasks'] for r in w.prepare())


def test_input_change_and_platform_switch_do_not_recreate(workspace):
    cfg, write = workspace
    write([row()])
    before = w.prepare()
    cfg['image_provider'] = 'mxapi'
    with pytest.raises(ValueError,match='冲突'):
        w.prepare()
    cfg['image_provider'] = 'tuzi'
    changed = row(subs=2)
    write([changed])
    with pytest.raises(ValueError,match='input_changed'):
        w.prepare()
    assert w.load_jsonl(w.batch_paths()['plan']) == before


def test_dry_run_no_network_no_batch_writes(workspace):
    _, write = workspace
    write([row(False,0)])
    paths = w.batch_paths()
    w.run_cycle(dry_run=True)
    assert not paths['root'].exists()


def test_url_object_dedup_and_column_numeric_order(workspace):
    _, write = workspace
    r = row(subs=1)
    r['已有副图链接2'] = r['已有副图链接1'] + '?token=123'
    r['已有副图链接10'] = f'{OSS}/SOURCE/new_sub10_SOURCE.png'
    r['已有副图链接3'] = f'{OSS}/SOURCE/new_sub3_SOURCE.png'
    write([r])
    planned = w.prepare()[0]
    assert planned['need_sub'] == 1
    urls = [i['url'] for i in planned['existing'] if i['role']=='sub']
    assert len(urls) == 3
    assert urls[1].endswith('new_sub3_SOURCE.png')
    assert urls[2].endswith('new_sub10_SOURCE.png')

@pytest.mark.parametrize('alter', ['missing_oss','reference_oss','role_conflict','different_base'])
def test_bad_inputs_block_without_generation(workspace, alter):
    _, write = workspace
    r = row()
    if alter == 'missing_oss':
        r['已有主图链接']=r['已有副图链接1']=''
    elif alter == 'reference_oss':
        r['参考主图链接']=r['已有主图链接']
    elif alter == 'role_conflict':
        r['已有副图链接1']=r['已有主图链接']
    else:
        r['已有副图链接1']=r['已有副图链接1'].replace('/walmart/','/different/')
    write([r])
    planned = w.prepare()[0]
    assert planned['errors'] and not planned['tasks']


def upload_rows(record, count=None):
    tasks=record['tasks'] if count is None else record['tasks'][:count]
    return [{'sku':record['record_id'],'image_name':t['image_name'],'status':'success',
        'oss_key':f"{record['oss_directory']['prefix']}/{t['image_name']}.png",
        'oss_url':f"{OSS}/SOURCE/{t['image_name']}.png"} for t in tasks]


def test_complete_three_outputs_and_resume_cumulative(workspace):
    _, write = workspace
    write([row(subs=1,old=5)])
    records = w.prepare()
    paths = w.batch_paths()
    w.save_jsonl(paths['sub_oss_checkpoint'],upload_rows(records[0]))
    payloads, manifest = w.build_results(records,paths)
    assert payloads[0]['complete']
    assert len(payloads[0]['secondary_urls'])==4
    assert '/EARLIEST/' in payloads[0]['main_image_url']
    assert all('/sub' not in url.rsplit('/',1)[-1] for url in payloads[0]['secondary_urls'])
    w.write_reports(payloads,manifest,paths)
    assert paths['latest_used_urls'].name == '最新使用图片URL.jsonl'
    assert paths['new_generated_urls'].name == '本批次新生成图片URL.jsonl'
    assert paths['archive_candidate_urls'].name == '待归档旧副图URL.jsonl'
    latest = w.load_jsonl(paths['latest_used_urls'])[0]
    new = w.load_jsonl(paths['new_generated_urls'])[0]
    archive = w.load_jsonl(paths['archive_candidate_urls'])[0]
    assert latest['complete'] and len(latest['secondary_urls'])==4
    assert new['main_image_url'] is None and len(new['secondary_urls'])==3
    assert len(archive['secondary_urls'])==5 and archive['archive_status']=='candidate_only'
    assert 'platform_update_status' not in latest
    assert 'update_main_image' not in latest
    wb = load_workbook(paths['latest_used_urls'].with_suffix('.xlsx'),read_only=True)
    assert wb.active.max_row==2
    assert '副图URL4' in next(wb.active.values)
    wb.close()
    second,_ = w.build_results(w.prepare(),paths)
    assert second == payloads


def test_partial_upload_preserves_old_main(workspace):
    _, write = workspace
    write([row(False,0,old=2)])
    records = w.prepare()
    paths = w.batch_paths()
    sub = [t for t in records[0]['tasks'] if t['role']=='sub']
    one={**records[0],'tasks':sub[:2]}
    w.save_jsonl(paths['sub_oss_checkpoint'],upload_rows(one))
    payloads,manifest = w.build_results(records,paths)
    assert not payloads[0]['complete']
    assert payloads[0]['main_image_url']==f'{OSS}/SOURCE/main_SOURCE.jpg'
    assert len(payloads[0]['secondary_urls'])==2
    assert len([i for i in manifest if i['kind']=='old'])==2


def test_all_complete_does_not_need_outputs_or_models(workspace,monkeypatch):
    _, write = workspace
    value=row(subs=6)
    value['参考主图链接']=value['参考副图链接1']=value['标题']=value['五点']=''
    write([value])
    w.prompt_stage(w.prepare(),w.batch_paths())
    monkeypatch.setattr(w,'image_stage',lambda *args:None)
    monkeypatch.setattr(w,'upload_stage',lambda *args:None)
    assert w.run_cycle()['complete']
    assert len(w.load_jsonl(w.batch_paths()['payload'])[0]['secondary_urls'])==6


def test_fixed_workbook_supports_unique_names_and_walmart_references(workspace):
    _, write=workspace
    write([row(False,0)])
    records=w.prepare()
    paths=w.batch_paths()
    plan={'image_plan':[{'image_number':i,'ai_image_generation_prompt':f'prompt{i}'} for i in range(1,5)]}
    w.save_json(paths['full_outputs']/f"{records[0]['record_id']}.json",plan)
    from ai_gateway.subtasks import mxapi_generate_images as engine
    for role in ('sub',):
        refs=w.build_image_input(records,paths,role)
        wb=load_workbook(paths[f'{role}_input'],read_only=True)
        values=list(wb.active.values)
        assert len(values)==(2 if role=='main' else 5)
        assert len(refs)==2
        assert all(v[1].startswith('https://i5.walmartimages.com/') for v in values[1:])
        assert all(v[5] for v in values[1:])
        wb.close()


def test_upload_uses_full_source_key_and_skips_own_success(workspace,monkeypatch,tmp_path):
    _,write=workspace
    write([row(False,0)])
    records=w.prepare()
    paths=w.batch_paths()
    task=next(t for t in records[0]['tasks'] if t['role']=='sub')
    image=tmp_path/'real.png'
    image.write_bytes(PNG)
    w.save_jsonl(paths['sub_checkpoint'],[{'sku':records[0]['record_id'],'image_name':task['image_name'],'status':'success',
        'downloaded_path':str(image),'file_size':len(PNG),'row_number':2}])
    from ai_gateway.clients import aliyun_oss_client as oss
    from ai_gateway.subtasks import oss_upload_images as uploader
    client=Mock()
    client.config=SimpleNamespace(max_retries=1)
    client.bucket.get_bucket_versioning.return_value.status=None
    client.public_url.side_effect=lambda key:f'https://test-bucket.oss-cn-beijing.aliyuncs.com/{key}'
    monkeypatch.setattr(oss,'load_aliyun_oss_config',lambda root:oss.AliyunOssConfig('key','secret','oss-cn-beijing.aliyuncs.com','test-bucket'))
    monkeypatch.setattr(oss,'AliyunOssClient',lambda cfg:client)
    w.upload_stage(records,paths,'sub')
    key=f"images/walmart/SOURCE/{task['image_name']}.png"
    client.bucket.put_object_from_file.assert_called_once_with(key,str(image))
    saved=w.load_jsonl(paths['sub_oss_checkpoint'])[0]
    assert saved['oss_key']==key
    assert '/images/images/' not in saved['oss_url']
    w.upload_stage(records,paths,'sub')
    assert client.bucket.put_object_from_file.call_count==1


def test_upload_error_does_not_become_success():
    error=RuntimeError('exists')
    error.status=409
    client=Mock()
    client.config=SimpleNamespace(max_retries=1)
    client.bucket.get_bucket_versioning.return_value.status=None
    client.bucket.put_object_from_file.side_effect=error
    result=w.PlanUploadClient(client).upload_file('unused','key')
    assert not result['success']
    assert client.bucket.put_object_from_file.call_count==1

@pytest.mark.parametrize('count',[1,2,3,4])
def test_dynamic_prompt_validation_and_default_six(count):
    from ai_gateway.subtasks.walmart_call_prompt_model import inspect_result_text, build_continue_payload
    text=json.dumps({'image_plan':[{'image_number':i,'ai_image_generation_prompt':f'concept {i}'} for i in range(1,count+1)]})
    assert w.validate_prompt(text,count)[1] is None
    assert inspect_result_text(text)[2] is not None
    payload=build_continue_payload({'metadata':{'expected_image_plan_count':count}}, {}, None, 'model', 100, None, 'auto','thinking')
    assert f'{count} 个对象' in payload['messages'][-1]['content']
    invalid=json.dumps({'image_plan':[{'image_number':1,'ai_image_generation_prompt':'x'}]*count})
    if count>1:
        assert w.validate_prompt(invalid,count)[1]


def test_generation_project_has_no_platform_update_code():
    task = ROOT / 'subtasks/walmart_image_replace'
    for name in ('06_replace_submit.py', '07_replace_reconcile.py', 'scripts/walmart_feed.py'):
        assert not (task / name).exists()
    cfg = json.loads((task / 'config.json').read_text(encoding='utf-8-sig'))
    assert 'walmart_api' not in cfg
    assert 'submit_replace' not in cfg['workflow'] and 'reconcile' not in cfg['workflow']


def test_old_batch_not_deleted(workspace):
    _,write=workspace
    write([row()])
    root=w.batch_paths()['root']
    root.mkdir(parents=True)
    old=root/'old_checkpoint.jsonl'
    old.write_text('old task id')
    with pytest.raises(ValueError,match='历史文件'):
        w.prepare()
    assert old.read_text()=='old task id'

def test_actual_shared_text_dynamic_results_and_recovery(workspace,monkeypatch):
    _,write=workspace
    write([row(subs=2)])
    records=w.prepare()
    paths=w.batch_paths()
    from ai_gateway.subtasks import walmart_call_prompt_model as shared
    from ai_gateway.config import loader
    from ai_gateway.config.loader import GatewayConfig
    from ai_gateway.clients import openai_chat_client
    gateway=GatewayConfig('tuzi_text','tuzi','https://example.test',api_key_value='test-only',max_retries=0)
    monkeypatch.setattr(loader,'load_app_config',lambda *args:SimpleNamespace(gateways={'tuzi_text':gateway}))
    fake=Mock()
    fake.gateway=gateway
    text=json.dumps({'image_plan':[{'image_number':i,'ai_image_generation_prompt':f'prompt {i}'} for i in (1,2)]})
    fake.responses_completions.return_value=({'id':'request','output_text':text},1)
    monkeypatch.setattr(openai_chat_client,'OpenAIChatClient',lambda gateway:fake)
    w.prompt_stage(records,paths)
    fake.responses_completions.assert_called_once()
    saved=w.load_jsonl(paths['model_results'])[0]
    assert saved['validation_status']=='passed'
    assert saved['image_plan_count']==2
    content=fake.responses_completions.call_args.args[0]['input'][0]['content']
    assert all('walmartimages.com' in i['image_url'] for i in content if i['type']=='input_image')
    canonical=paths['full_outputs']/f"{records[0]['record_id']}.json"
    canonical.unlink()  # Simulate interruption after raw model output, before canonical checkpoint.
    paths['model_results'].unlink()
    w.prompt_stage(records,paths)
    assert canonical.exists()
    assert fake.responses_completions.call_count==1


def test_shared_image_engine_submit_query_429_resume_without_new_tasks(workspace,monkeypatch):
    _,write=workspace
    write([row(subs=3)])
    records=w.prepare()
    paths=w.batch_paths()
    w.save_json(paths['full_outputs']/f"{records[0]['record_id']}.json",{'image_plan':[{'image_number':1,'ai_image_generation_prompt':'scene'}]})
    from ai_gateway.subtasks import mxapi_generate_images as engine
    from ai_gateway.clients.image_providers import TuziImageAdapter
    from ai_gateway.config.loader import GatewayConfig
    gateway=GatewayConfig('tuzi','tuzi','https://example.test',api_key_value='offline',max_retries=0)
    adapter=TuziImageAdapter(gateway)
    submit=Mock(return_value=({'id':'task-one'},1))
    query=Mock(return_value=({'status':'completed','video_url':'https://image.test/output'},1))
    downloads=[]
    def download(url,path,**kwargs):
        downloads.append(url)
        if len(downloads)==1:
            raise RuntimeError('HTTP 429: Too Many Requests')
        path.parent.mkdir(parents=True,exist_ok=True)
        path.write_bytes(PNG)
        return len(PNG)
    monkeypatch.setattr(adapter,'submit',submit)
    monkeypatch.setattr(adapter,'query',query)
    monkeypatch.setattr(adapter,'download',download)
    monkeypatch.setattr(engine,'load_app_config',lambda *args:SimpleNamespace(gateways={'tuzi':gateway}))
    monkeypatch.setattr(engine,'create_image_adapter',lambda *args:adapter)
    original=w.load_stage_config
    def stage(name):
        data=original(name)
        if name in ('generate_images','generate_main_image'):
            data['limits']['submit_delay_seconds']=0
            data['retry'].update(retry_delay_seconds=0,query_retry_delay_seconds=0)
        return data
    monkeypatch.setattr(w,'load_stage_config',stage)
    w.image_stage(records,paths,'sub')
    first=w.load_jsonl(paths['sub_checkpoint'])[-1]
    assert first['task_id']=='task-one' and first['status']=='submitted'
    w.image_stage(records,paths,'sub')
    second=w.load_jsonl(paths['sub_checkpoint'])[-1]
    assert second['task_id']=='task-one' and second['status']=='pending'
    assert submit.call_count==1 and len(downloads)==1
    w.image_stage(records,paths,'sub')
    third=w.load_jsonl(paths['sub_checkpoint'])[-1]
    assert third['task_id']=='task-one' and third['status']=='success'
    assert third['attempts']==0
    assert submit.call_count==1 and len(downloads)==2
    w.image_stage(records,paths,'sub')
    assert submit.call_count==1 and len(downloads)==2


def test_remote_uploaded_image_not_regenerated_when_local_file_lost(workspace):
    _,write=workspace
    write([row(subs=3)])
    records=w.prepare()
    paths=w.batch_paths()
    w.save_json(paths['full_outputs']/f"{records[0]['record_id']}.json",{'image_plan':[{'image_number':1,'ai_image_generation_prompt':'scene'}]})
    w.save_jsonl(paths['sub_oss_checkpoint'],upload_rows(records[0]))
    assert w.input_rows(records,paths,'sub')==[]
    assert w.build_results(records,paths)[0][0]['complete']


def test_record_limit_same_across_main_and_sub(workspace):
    cfg,write=workspace
    cfg['execution']['max_records']=1
    write([row(False,0,result='B'),row(False,0,result='C')])
    records=w.prepare()
    assert len(records)==1 and records[0]['result_sku']=='B'
    assert len(records[0]['tasks'])==4


@pytest.mark.parametrize('versioning', ['Enabled', 'Suspended', None])
def test_versioned_bucket_allows_repeated_planned_upload(tmp_path, versioning):
    image = tmp_path / 'generated.png'
    image.write_bytes(PNG)
    client = Mock()
    client.config = SimpleNamespace(max_retries=1)
    client.bucket.get_bucket_versioning.return_value.status = versioning
    uploader = w.PlanUploadClient(client)
    key = 'images/walmart/SOURCE/new_sub_SOURCE_unique.png'
    assert uploader.upload_file(image, key)['success']
    # Simulate upload success followed by a crash before the checkpoint was saved.
    assert uploader.upload_file(image, key)['success']
    assert client.bucket.put_object_from_file.call_count == 2
    client.bucket.put_object_from_file.assert_called_with(key, str(image))
    client.bucket.get_bucket_versioning.assert_not_called()
    client.bucket.object_exists.assert_not_called()

def test_ready_cross_base_images_do_not_require_upload_directory(workspace):
    _,write=workspace
    value=row(subs=4)
    value['已有主图链接']=value['已有主图链接'].replace('/walmart/','/historical_main/')
    value['参考主图链接']=value['参考副图链接1']=value['标题']=value['五点']=''
    write([value])
    records=w.prepare()
    assert not records[0]['errors'] and not records[0]['tasks']
    assert records[0]['oss_directory'] is None
    assert w.build_results(records,w.batch_paths())[0][0]['complete']


def test_old_main_is_retained_without_generation_or_archive(workspace, monkeypatch):
    _, write = workspace
    value = row(False, 4)
    value['已有主图链接'] = REF
    value['参考主图链接'] = value['标题'] = value['五点'] = ''
    write([value])
    records = w.prepare()
    assert not records[0]['tasks']
    monkeypatch.setattr(w, 'image_stage', lambda *args: None)
    monkeypatch.setattr(w, 'upload_stage', lambda *args: None)
    assert w.run_cycle()['complete']
    paths = w.batch_paths()
    assert w.load_jsonl(paths['latest_used_urls'])[0]['main_image_url'] == REF
    assert w.load_jsonl(paths['archive_candidate_urls'])[0]['main_image_url'] is None
    assert w.load_jsonl(paths['new_generated_urls'])[0]['main_image_url'] is None


def test_main_remote_helpers_are_disabled(workspace):
    _, write = workspace
    write([row(False, 4)])
    records = w.prepare()
    with pytest.raises(ValueError, match='禁止生成主图'):
        w.image_stage(records, w.batch_paths(), 'main')
    with pytest.raises(ValueError, match='禁止上传主图'):
        w.upload_stage(records, w.batch_paths(), 'main')


@pytest.mark.parametrize('status,reason', [
    ('failed', 'HTTP 403: temporary service unavailable'),
    ('failed', 'HTTP 429: exceeded rate limit'),
    ('invalid', '提示词编号必须按计划为1..need_sub'),
])
def test_prompt_failure_in_payload_and_all_reports(workspace, status, reason):
    _, write = workspace
    write([row(result='FAILED'), row(result='OTHER')])
    records = w.prepare()
    paths = w.batch_paths()
    w.save_jsonl(paths['model_results'], [{'sku': records[0]['record_id'], 'status': status, 'error_message': reason}])
    payloads, manifests = w.build_results(records, paths)
    expected = f'generate_prompts: {reason}'
    assert expected in payloads[0]['errors']
    assert expected not in payloads[1]['errors']
    assert not payloads[0]['complete']
    w.write_reports(payloads, manifests, paths)
    for name in ('latest_used_urls', 'new_generated_urls', 'archive_candidate_urls'):
        assert expected in w.load_jsonl(paths[name])[0]['errors']
        book = load_workbook(paths[name].with_suffix('.xlsx'), read_only=True)
        values = list(book.active.values)
        error_column = values[0].index('失败原因')
        assert expected in values[1][error_column]
        assert expected not in values[2][error_column]
        book.close()
    # Next round succeeds: the old failure must disappear from derived outputs.
    w.save_jsonl(paths['model_results'], [{'sku': records[0]['record_id'], 'status': 'success'}])
    assert expected not in w.build_results(records, paths)[0][0]['errors']


def test_recovered_prompt_and_completed_uploads_hide_stale_failure(workspace):
    _, write = workspace
    write([row()])
    records = w.prepare()
    paths = w.batch_paths()
    reason = 'HTTP 429: exceeded rate limit'
    w.save_jsonl(paths['model_results'], [{'sku': records[0]['record_id'], 'status': 'failed', 'error_message': reason}])
    output = paths['full_outputs'] / f"{records[0]['record_id']}.json"
    w.save_json(output, {'image_plan': [{'image_number': i, 'ai_image_generation_prompt': f'concept{i}'} for i in range(1, 4)]})
    assert f'generate_prompts: {reason}' not in w.build_results(records, paths)[0][0]['errors']
    output.unlink()
    w.save_jsonl(paths['sub_oss_checkpoint'], upload_rows(records[0]))
    payloads, _ = w.build_results(records, paths)
    assert payloads[0]['complete'] and not payloads[0]['errors']



def test_manifest_input_columns_and_recorded_timestamps(workspace):
    _, write = workspace
    value = row(subs=0)
    value['已有副图链接2'] = f'{OSS}/SOURCE/new_sub2_SOURCE.png'
    value['已有副图链接10'] = f'{OSS}/SOURCE/sub10_SOURCE.jpg'
    value['已有副图链接11'] = value['已有副图链接2'] + '?duplicate=1'
    write([value])
    records = w.prepare()
    paths = w.batch_paths()
    task = records[0]['tasks'][0]
    planned_at = task['plan_created_at']
    assert planned_at.endswith('+00:00')
    assert w.prepare()[0]['tasks'][0]['plan_created_at'] == planned_at
    uploads = upload_rows(records[0], 1)
    uploads[0]['created_at'] = '2026-09-18T19:01:02'
    w.save_jsonl(paths['sub_oss_checkpoint'], uploads)
    w.save_jsonl(paths['sub_checkpoint'], [{'sku': records[0]['record_id'], 'image_name': task['image_name'],
        'status': 'success', 'created_at': '2026-09-18T19:00:01'}])
    _, manifest = w.build_results(records, paths)
    existing = [i for i in manifest if i['origin'] == 'existing']
    assert [(i['input_column'], i['input_order']) for i in existing] == [
        ('已有主图链接', 1), ('已有副图链接2', 2), ('已有副图链接10', 10)]
    assert all(i['plan_created_at'] is None and i['generation_recorded_at'] is None for i in existing)
    generated = next(i for i in manifest if i.get('image_id') == task['image_id'])
    assert generated['plan_created_at'] == planned_at
    assert generated['generation_recorded_at'] == '2026-09-18T19:00:01'
    assert generated['upload_recorded_at'] == '2026-09-18T19:01:02'
    assert generated['ordinal'] == 1 and generated['input_column'] is None
    assert generated['source_sku'] == 'SOURCE' and generated['result_sku'] == 'RESULT'
    assert w.build_results(w.prepare(), paths)[1] == manifest


def test_manifest_old_plan_enrichment_does_not_rewrite_history(workspace):
    _, write = workspace
    write([row()])
    records = w.prepare()
    paths = w.batch_paths()
    for image in records[0]['existing']:
        image.pop('input_column')
        image.pop('input_order')
    for task in records[0]['tasks']:
        task.pop('plan_created_at')
    w.save_jsonl(paths['plan'], records)
    before = paths['plan'].read_bytes()
    _, manifest = w.build_results(w.prepare(), paths)
    assert paths['plan'].read_bytes() == before
    assert manifest[0]['input_column'] == '已有主图链接'
    assert manifest[1]['input_column'] == '已有副图链接1'
    for image in manifest:
        assert image['plan_created_at'] is None
        assert image['generation_recorded_at'] is None
        assert image['upload_recorded_at'] is None


def test_product_metadata_in_image_set_and_report(workspace):
    _, write = workspace
    value = row(subs=4)
    value['GTIN'] = '00012345678905'
    value['商品类型'] = 'Hair Brushes'
    write([value])
    records = w.prepare()
    paths = w.batch_paths()
    payloads, manifests = w.build_results(records, paths)
    assert payloads[0]['product_type'] == 'Hair Brushes'
    assert payloads[0]['gtin'] == '00012345678905'
    assert paths['payload'].name == 'image_set.jsonl'
    w.write_reports(payloads, manifests, paths)
    book = load_workbook(paths['latest_used_urls'].with_suffix('.xlsx'), read_only=True)
    values = list(book.active.values)
    assert values[1][values[0].index('GTIN')] == '00012345678905'
    assert values[1][values[0].index('商品类型')] == 'Hair Brushes'
    assert '平台更新状态' not in values[0]
    book.close()



def load_review_module(monkeypatch):
    spec = importlib.util.spec_from_file_location('replace_review_test', ROOT / 'subtasks/walmart_image_replace/09_replace_export_review.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    monkeypatch.setattr(module, 'workflow', w)
    return module


def review_png():
    import io
    from PIL import Image
    buffer = io.BytesIO()
    Image.new('RGB', (32, 20), 'red').save(buffer, format='PNG')
    return buffer.getvalue()


def test_review_groups_images_local_generation_cache_and_human_fields(workspace, monkeypatch, tmp_path):
    _, write = workspace
    value = row(subs=1)
    value['GTIN'] = '00012345678905'
    value['标题'] = '=unsafe formula'
    write([value])
    source = Path(w.load_task_config()['input']['excel_path'])
    source_book = load_workbook(source)
    source_header = [cell.value for cell in source_book.active[1]]
    source_book.active.cell(2, source_header.index('标题') + 1).data_type = 's'
    source_book.save(source)
    source_book.close()
    records = w.prepare()
    paths = w.batch_paths()
    generated = []
    for index, task in enumerate(records[0]['tasks']):
        image = tmp_path / f'generated{index}.png'
        image.write_bytes(review_png())
        generated.append({'sku': records[0]['record_id'], 'image_name': task['image_name'], 'status': 'success', 'downloaded_path': str(image)})
    w.save_jsonl(paths['sub_checkpoint'], generated)
    w.save_jsonl(paths['sub_oss_checkpoint'], upload_rows(records[0]))
    before = paths['plan'].read_bytes()
    review = load_review_module(monkeypatch)
    downloader = Mock(return_value=review_png())
    monkeypatch.setattr(review, 'download_image', downloader)
    output = tmp_path / '人工审核预览.xlsx'
    summary = review.export(output=output)
    assert summary['records'] == summary['complete'] == 1
    assert summary['embedded_images'] == 7 and summary['preview_errors'] == 0
    assert downloader.call_count == 4  # Two references, retained main, existing new sub; generated images use local files.
    assert paths['plan'].read_bytes() == before
    book = load_workbook(output)
    sheet = book['人工审核预览']
    header = [cell.value for cell in sheet[2]]
    assert sheet.cell(1, 13).value == '参考图（沃尔玛平台素材）'
    assert sheet.cell(1, 15).value == '最新使用图（原主图＋最终副图）'
    assert sheet.cell(2, 13).fill.fgColor.rgb != sheet.cell(2, 15).fill.fgColor.rgb
    assert sheet.cell(3, header.index('GTIN') + 1).value == '00012345678905'
    assert '描述' not in header
    title = sheet.cell(3, header.index('标题') + 1)
    assert title.value == '=unsafe formula' and title.data_type == 's'
    assert sheet.cell(3, header.index('审核结论') + 1).value == '待审核'
    assert sheet.cell(4, 13).hyperlink.target == REF
    assert sheet.cell(4, 15).hyperlink.target.endswith('new_main2_EARLIEST.png')
    assert len(sheet._images) == 7
    assert {image.anchor._from.col for image in sheet._images} == set(range(12, 19))
    book.close()
    downloader.reset_mock()
    cached_output = tmp_path / '离线审核.xlsx'
    offline_summary = review.export(output=cached_output, offline=True)
    assert offline_summary['embedded_images'] == 3
    assert offline_summary['preview_errors'] == 4
    downloader.assert_not_called()
    with pytest.raises(ValueError, match='保护人工意见'):
        review.export(output=output)


def test_review_partial_results_failures_and_dry_run_do_not_hide_records(workspace, monkeypatch, tmp_path):
    _, write = workspace
    write([row(False, 1, result='A'), row(False, 5, result='B')])
    w.prepare()
    review = load_review_module(monkeypatch)
    downloader = Mock(side_effect=RuntimeError('HTTP 429 preview download'))
    monkeypatch.setattr(review, 'download_image', downloader)
    output = tmp_path / '审核.xlsx'
    summary = review.export(output=output, dry_run=True, max_skus_per_file=1)
    assert summary['records'] == 2 and summary['complete'] == 1
    assert not output.exists() and not (w.batch_paths()['root'] / '09_review').exists()
    downloader.assert_not_called()
    result = review.export(output=output, offline=True, max_skus_per_file=1)
    assert len(result['files']) == 2 and result['preview_errors'] > 0
    downloader.assert_not_called()
    first = load_workbook(result['files'][0])
    sheet = first['人工审核预览']
    header = [cell.value for cell in sheet[2]]
    assert sheet.cell(3, header.index('图集是否完整') + 1).value is False
    assert '不足4张' in sheet.cell(3, header.index('失败原因') + 1).value
    assert '离线模式' in sheet.cell(3, header.index('预览加载问题') + 1).value
    assert sheet.cell(4, header.index('最新使用主图（原样保留）') + 1).hyperlink.target.endswith('main_SOURCE.jpg')
    assert all('sub2_SOURCE.jpg' not in str(cell.value) for row_cells in sheet for cell in row_cells)
    first.close()


def test_append_rows_preserves_paid_plan_and_checkpoints(workspace):
    _, write = workspace
    write([row(result='A')])
    original = w.prepare()
    paths = w.batch_paths()
    checkpoint = [{'sku': original[0]['record_id'], 'image_name': original[0]['tasks'][0]['image_name'], 'task_id': 'remote-paid', 'status': 'pending'}]
    w.save_jsonl(paths['sub_checkpoint'], checkpoint)
    plan_before = paths['plan'].read_bytes()
    checkpoint_before = paths['sub_checkpoint'].read_bytes()
    write([row(result='A'), row(result='B')])
    preview = w.prepare(dry_run=True)
    assert len(preview) == 2 and preview[0] == original[0]
    assert paths['plan'].read_bytes() == plan_before
    assert paths['sub_checkpoint'].read_bytes() == checkpoint_before
    appended = w.prepare()
    assert len(appended) == 2 and appended[0] == original[0]
    assert appended[1]['row_number'] == 3 and len(appended[1]['tasks']) == 3
    assert w.prepare() == appended
    assert paths['sub_checkpoint'].read_bytes() == checkpoint_before
    assert w.load_jsonl(paths['inventory']) == appended
    names = [t['image_name'] for r in appended for t in r['tasks']]
    assert len(names) == len(set(names))


@pytest.mark.parametrize('change', ['delete', 'insert', 'reorder', 'duplicate'])
def test_append_rejects_changes_to_existing_rows(workspace, change):
    _, write = workspace
    rows = [row(result='A'), row(result='B')]
    write(rows)
    original = w.prepare()
    changed = {'delete': rows[:1], 'insert': [row(result='C'), *rows],
               'reorder': list(reversed(rows)), 'duplicate': [*rows, row(result='A')]}[change]
    write(changed)
    with pytest.raises(ValueError, match='input_changed'):
        w.prepare()
    assert w.load_jsonl(w.batch_paths()['plan']) == original


def test_append_by_increasing_record_limit(workspace):
    cfg, write = workspace
    write([row(result='A'), row(result='B')])
    cfg['execution']['max_records'] = 1
    original = w.prepare()
    cfg['execution']['max_records'] = 2
    appended = w.prepare()
    assert len(appended) == 2 and appended[0] == original[0]
    cfg['execution']['max_records'] = 1
    with pytest.raises(ValueError, match='input_changed'):
        w.prepare()
    assert w.load_jsonl(w.batch_paths()['plan']) == appended
