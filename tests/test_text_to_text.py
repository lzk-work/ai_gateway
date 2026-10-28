from pathlib import Path
import sys

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from ai_gateway.text_to_text.engine import TextBatchEngine
from ai_gateway.text_to_text.config import TaskDefinition, load_task
from ai_gateway.text_to_text.models import CallResult
from ai_gateway.text_to_text.response import parse_response
from ai_gateway.text_to_text.state_store import StateStore
from ai_gateway.text_to_text.template import PromptTemplate, render_inline, request_hash
from ai_gateway.text_to_text.tuzi import TuziTextGateway


def test_headers_render_directly_as_placeholders():
    template = PromptTemplate("标题：{{ Product Title }}\n描述：{{ 商品描述 }}")
    template.validate_headers(["Product Title", "商品描述"])
    assert template.render({"Product Title": "Lamp", "商品描述": "台灯"}) == "标题：Lamp\n描述：台灯"


def test_placeholder_matching_is_case_insensitive():
    assert PromptTemplate("{{ SKU }}").render({"sku": "A001"}) == "A001"


def test_missing_placeholder_header_fails_before_api():
    with pytest.raises(ValueError, match="没有对应表头"):
        PromptTemplate("{{ missing }}").validate_headers(["sku", "title"])


def test_inline_record_key_supports_composite_fields():
    assert render_inline("{{ sku }}::{{ language }}", {"SKU": "1", "Language": "zh"}) == "1::zh"


def test_request_hash_changes_with_prompt_or_contract():
    first = request_hash("A", {"model": "m1"})
    assert first != request_hash("B", {"model": "m1"})
    assert first != request_hash("A", {"model": "m2"})


def test_fixed_fields_parser():
    parsed = parse_response("SKU: A1\n产品识别: 台灯\n提取关键词: 灯具", {
        "type": "fixed_fields", "fields": ["SKU", "产品识别", "提取关键词"]
    })
    assert parsed == {"SKU": "A1", "产品识别": "台灯", "提取关键词": "灯具"}


def test_fixed_fields_parser_accepts_numbered_markdown_bold_labels():
    parsed = parse_response(
        "1. **SKU**：A1\n2、**产品识别**：台灯\n3—**总体风险等级**：LOW",
        {
            "type": "fixed_fields",
            "fields": ["SKU", "产品识别", "总体风险等级"],
        },
    )
    assert parsed == {"SKU": "A1", "产品识别": "台灯", "总体风险等级": "LOW"}


def test_state_store_is_authoritative():
    from uuid import uuid4

    path = ROOT / "tests" / f".text_to_text_state_{uuid4().hex}.sqlite3"
    store = StateStore(path)
    try:
        result = CallResult("A1", 2, "hash", "success", "ok", {"answer": "x"}, "model", "req", 10, 1, updated_at="now")
        store.save(result)
        saved = store.get("A1")
        assert saved["status"] == "success"
        assert saved["parsed"] == {"answer": "x"}
    finally:
        store.close()
        path.unlink(missing_ok=True)
        Path(str(path) + "-wal").unlink(missing_ok=True)
        Path(str(path) + "-shm").unlink(missing_ok=True)


def test_stage_changed_input_resets_old_success_without_attempt():
    from ai_gateway.text_to_text.models import InputRecord
    from uuid import uuid4

    path = ROOT / "tests" / f".text_to_text_stage_{uuid4().hex}.sqlite3"
    store = StateStore(path)
    try:
        store.save(CallResult("A1", 2, "old", "success", "old text", {}, updated_at="old"))
        store.stage(InputRecord("A1", 2, {}, "new prompt", "new"))
        saved = store.get("A1")
        assert saved["status"] == "pending"
        assert saved["input_hash"] == "new"
        assert saved["raw_text"] == ""
        assert saved["attempts"] == 0
    finally:
        store.close()
        path.unlink(missing_ok=True)
        Path(str(path) + "-wal").unlink(missing_ok=True)
        Path(str(path) + "-shm").unlink(missing_ok=True)


def test_resume_decision_uses_status_and_input_hash():
    assert TextBatchEngine._should_run(None, "new")
    assert not TextBatchEngine._should_run({"status": "success", "input_hash": "same"}, "same")
    assert TextBatchEngine._should_run({"status": "success", "input_hash": "old"}, "new")
    assert TextBatchEngine._should_run({"status": "api_transient_failed", "input_hash": "same"}, "same")
    assert not TextBatchEngine._should_run({"status": "input_invalid", "input_hash": "same"}, "same")


def test_retry_failed_selects_all_failures_but_not_success():
    for status in ("api_transient_failed", "api_permanent_failed", "validation_failed"):
        assert TextBatchEngine._should_run(
            {"status": status, "input_hash": "same"},
            "same",
            retry_failed=True,
        )
    assert not TextBatchEngine._should_run(
        {"status": "success", "input_hash": "same"},
        "same",
        retry_failed=True,
    )


def test_generic_raw_text_task_dry_run_with_synthetic_excel():
    from openpyxl import Workbook
    from uuid import uuid4

    path = ROOT / "tests" / f".text_to_text_input_{uuid4().hex}.xlsx"
    try:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["record_id", "text", "处理动作"])
        sheet.append(["R1", "hello", "自动"])
        sheet.append(["R1", "world", "自动"])
        workbook.save(path)
        task_root = ROOT / "subtasks" / "text_to_text"
        task = TaskDefinition(
            "synthetic_raw_text",
            "1.0.0",
            ROOT / "tests" / "synthetic_raw_text",
            {
                "input": {"type": "excel", "header_row": 1},
                "prompt": {"template": "prompt.txt"},
                "response": {"type": "raw_text"},
            },
            "{{ text }}",
        )
        engine = TextBatchEngine(ROOT, task_root, task, "synthetic_dry_run")
        summary = engine.dry_run(path)
        assert summary["records"] == 2
        assert summary["placeholders"] == ["text"]
        _, records, _ = engine.prepare(path, create_snapshot=False)
        assert [record.record_key for record in records] == ["row:2", "row:3"]
        assert not engine.paths.root.exists()
    finally:
        path.unlink(missing_ok=True)


def test_batch_output_is_nested_under_its_task():
    task_root = ROOT / "subtasks" / "text_to_text"
    task = load_task(task_root / "tasks", "keyword_1688")
    engine = TextBatchEngine(ROOT, task_root, task, "batch_001")
    assert engine.paths.root == task.root / "batches" / "batch_001"


def test_multiple_image_columns_and_multiline_cells_are_collected_once():
    fields = {
        "Main Image": "https://example.com/a.jpg",
        "Other Images": "https://example.com/b.jpg\nhttps://example.com/a.jpg",
    }
    assert TextBatchEngine._image_urls(fields, ["Main Image", "Other Images"]) == [
        "https://example.com/a.jpg",
        "https://example.com/b.jpg",
    ]


def test_zero_image_placeholder_is_ignored():
    assert TextBatchEngine._image_urls(
        {"Main Image": 0, "Other Images": "0\nhttps://example.com/b.jpg"},
        ["Main Image", "Other Images"],
    ) == ["https://example.com/b.jpg"]


def test_required_main_image_is_rejected_before_api():
    from openpyxl import Workbook
    from uuid import uuid4

    path = ROOT / "tests" / f".patent_review_input_{uuid4().hex}.xlsx"
    try:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "SKU", "Product Title", "Bullet Points", "Description", "1688 Title",
            "主图url", "副图url1", "副图url2", "副图url3", "副图url4", "副图url5", "副图url6",
        ])
        sheet.append(["A1", "title", "points", "description", "1688", 0, "https://example.com/sub.jpg"])
        workbook.save(path)
        task_root = ROOT / "subtasks" / "text_to_text"
        task = load_task(task_root / "tasks", "patent_review")
        engine = TextBatchEngine(ROOT, task_root, task, "required_main_image_test")
        _, records, _ = engine.prepare(path, create_snapshot=False)
        assert records[0].fields["__missing_required__"] == ["主图url"]
        assert records[0].image_urls == ["https://example.com/sub.jpg"]
        summary = engine.dry_run(path)
        assert summary["input_invalid"] == 1
        assert summary["input_invalid_rows"] == [2]
    finally:
        path.unlink(missing_ok=True)


def test_tuzi_gateway_uses_multimodal_content_for_images():
    gateway = object.__new__(TuziTextGateway)

    class FakeClient:
        def __init__(self):
            self.payload = None

        def chat_completions(self, payload):
            self.payload = payload
            return {"id": "req-1", "choices": [{"message": {"content": "ok"}}]}, 12

    gateway.client = FakeClient()
    gateway.call(
        prompt="describe",
        model="vision-model",
        max_tokens=100,
        temperature=0.1,
        image_urls=["https://example.com/a.jpg", "https://example.com/b.jpg"],
        image_detail="high",
    )
    content = gateway.client.payload["messages"][0]["content"]
    assert content == [
        {"type": "text", "text": "describe"},
        {"type": "image_url", "image_url": {"url": "https://example.com/a.jpg", "detail": "high"}},
        {"type": "image_url", "image_url": {"url": "https://example.com/b.jpg", "detail": "high"}},
    ]


def test_tuzi_gateway_keeps_plain_text_content_without_images():
    gateway = object.__new__(TuziTextGateway)

    class FakeClient:
        def __init__(self):
            self.payload = None

        def chat_completions(self, payload):
            self.payload = payload
            return {"id": "req-1", "choices": [{"message": {"content": "ok"}}]}, 12

    gateway.client = FakeClient()
    gateway.call(prompt="plain", model="text-model", max_tokens=100, temperature=0.1)
    assert gateway.client.payload["messages"][0]["content"] == "plain"
