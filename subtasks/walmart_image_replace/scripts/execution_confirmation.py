"""Execution confirmation summary for the Walmart image-replace workflow.

执行确认视图：批次状态、阶段开关、各阶段产物状态。只读，不调任何外部接口。
06/07 沃尔玛 Feed 提交与对账已接入真实链路（真实运行会调用沃尔玛 API）。
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from workflow_common import (  # noqa: E402
    batch_name_from_input,
    batch_paths,
    load_jsonl,
    load_task_config,
    task_execution,
    task_input,
    walmart_stores,
    workflow_switches,
)


def count_input_rows() -> int:
    from openpyxl import load_workbook

    excel_path = task_input().get("excel_path")
    if not excel_path or not Path(excel_path).exists():
        return 0
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb[task_input().get("sheet_name")] if task_input().get("sheet_name") else wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return sum(1 for r in rows[1:] if any(c is not None and str(c).strip() for c in r))


def main() -> None:
    task_config = load_task_config()
    execution = task_execution()
    switches = workflow_switches()
    paths = batch_paths()
    batch_name = batch_name_from_input()
    input_cfg = task_input()
    stores = walmart_stores()

    print("\n=== 执行前确认（walmart_image_replace）===")
    print(f"业务任务: {task_config.get('name', 'walmart_image_replace')}")
    print(f"批次名: {batch_name}")
    print(f"批次目录: {paths['root']}")
    print(f"入参Excel: {input_cfg.get('excel_path', '')}")
    print(f"Sheet: {input_cfg.get('sheet_name') or '(active)'}")
    print(f"Excel有效行数: {count_input_rows()} | 业务 max_records: {execution.get('max_records')}")
    print(
        "阶段开关: "
        f"01={switches['query_existing_images']} | 02={switches['generate_prompts']} | "
        f"03={switches['generate_images']} | 04={switches['upload_oss']} | "
        f"05={switches['build_replace_result']} | 06={switches['submit_replace']} | "
        f"07={switches['reconcile']} | 08={switches['build_report']}"
    )
    print(f"店铺配置: {', '.join(stores.keys()) or '无'}")

    # 各阶段产物状态
    stage_files = [
        ("01 image_inventory", paths["inventory"]),
        ("02 model_results", paths["model_results"]),
        ("03 image_input", paths["image_input_excel"]),
        ("03 image_generation_results", paths["image_results"]),
        ("04 oss_upload_results", paths["oss_results"]),
        ("05 replace_payload", paths["payload"]),
        ("06 submit_results", paths["submit_results"]),
        ("07 reconcile_results", paths["reconcile_results"]),
        ("08 result_report", paths["report"]),
    ]
    print("\n--- 批次产物状态 ---")
    for label, path in stage_files:
        if path.suffix == ".jsonl":
            n = len(load_jsonl(path)) if path.exists() else 0
            print(f"{label}: {'缺失' if not path.exists() else f'{n} 条'} ({path.name})")
        else:
            print(f"{label}: {'缺失' if not path.exists() else '存在'} ({path.name})")

    # 可提交 SKU 概览（替换结果完整性 + 06/07 提交对账状态）
    payloads = load_jsonl(paths["payload"])
    if payloads:
        eligible = [p for p in payloads if p.get("complete") and p["store"] in stores and p.get("main_image_url")]
        print("\n--- 替换结果概览 ---")
        print(f"payload 总数: {len(payloads)} | 完整(complete+store+主图): {len(eligible)}")
        if eligible:
            by_store: dict[str, list] = {}
            for p in eligible:
                by_store.setdefault(p["store"], []).append(p["platform_sku"])
            for store, skus in by_store.items():
                print(f"  {store}: {len(skus)} 个 SKU -> {[s for s in skus]}")
        submits = load_jsonl(paths["submit_results"])
        reconciles = load_jsonl(paths["reconcile_results"])
        if submits:
            ok = sum(1 for s in submits if s.get("status") in ("submitted", "success", "processed"))
            feeds = {s.get("feed_id") for s in submits if s.get("feed_id")}
            print(f"[06 提交] {len(submits)} 行（成功提交 {ok}）| feedId: {sorted(feeds)}")
        else:
            print("[06 提交] 未运行")
        if reconciles:
            from collections import Counter as _C
            dist = dict(_C(r.get("status") for r in reconciles))
            print(f"[07 对账] {len(reconciles)} 行，状态分布: {dist}")
        else:
            print("[07 对账] 未运行")


if __name__ == "__main__":
    main()
