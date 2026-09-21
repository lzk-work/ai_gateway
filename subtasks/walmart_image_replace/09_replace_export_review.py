"""Independent human review export: reference images versus final image sets."""
from __future__ import annotations
import argparse
from concurrent.futures import ThreadPoolExecutor, as_completed
import io
import json
from pathlib import Path
import sys
import threading

TASK_ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(TASK_ROOT))
import replace_workflow_common as workflow
import requests
from PIL import Image, ImageOps
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from ai_gateway.subtasks.walmart_call_prompt_model import sanitize_excel_value

INFO_HEADERS = ['来源SKU', '结果SKU', '店铺', 'GTIN', '商品类型', '标题', '五点',
                '图集是否完整', '失败原因', '审核结论', '审核意见', '预览加载问题']
REFERENCE_COLOR = 'DDEBF7'
FINAL_COLOR = 'E2F0D9'
MAX_IMAGE_BYTES = 20 * 1024 * 1024


def thumbnail(content, size):
    with Image.open(io.BytesIO(content)) as source:
        source.load()
        image = ImageOps.exif_transpose(source).convert('RGBA')
        image.thumbnail((size, size))
        background = Image.new('RGB', image.size, 'white')
        background.paste(image, mask=image.getchannel('A'))
        buffer = io.BytesIO()
        background.save(buffer, format='JPEG', quality=75)
        return buffer.getvalue(), image.width, image.height


def download_image(url, timeout=30):
    workflow.url_parts(url)
    with requests.get(url, timeout=timeout, stream=True) as response:
        response.raise_for_status()
        content = bytearray()
        for chunk in response.iter_content(64 * 1024):
            content.extend(chunk)
            if len(content) > MAX_IMAGE_BYTES:
                raise ValueError('图片超过20MB预览下载上限')
        return bytes(content)


class PreviewImages:
    def __init__(self, size, offline=False, concurrency=8):
        self.size, self.offline, self.concurrency = size, offline, concurrency
        self.memo = {}
        self.lock = threading.Lock()

    def get(self, url, local_path=None):
        if url in self.memo:
            result = self.memo[url]
            if isinstance(result, Exception):
                raise result
            return result
        local_error = None
        if local_path:
            path = Path(local_path)
            if not path.is_absolute():
                path = workflow.PROJECT_ROOT / path
            if path.is_file():
                try:
                    result = thumbnail(path.read_bytes(), self.size)
                    self.memo[url] = result
                    return result
                except Exception as exc:
                    local_error = str(exc)
        if self.offline:
            raise ValueError('离线模式：无可用本地图片' + (f'；{local_error}' if local_error else ''))
        result = thumbnail(download_image(url), self.size)
        with self.lock:
            self.memo[url] = result
        return result

    def preload(self, items):
        unique = {}
        for url, local_path in items:
            if url and url not in self.memo:
                unique.setdefault(url, local_path)
        with ThreadPoolExecutor(max_workers=self.concurrency) as pool:
            futures = {pool.submit(self.get, url, path): url for url, path in unique.items()}
            for future in as_completed(futures):
                try:
                    future.result()
                except Exception as exc:
                    with self.lock:
                        self.memo[futures[future]] = exc


def set_cell(sheet, row, column, value):
    cell = sheet.cell(row, column, sanitize_excel_value(value))
    if isinstance(cell.value, str):
        cell.data_type = 's'
    cell.alignment = Alignment(vertical='top', wrap_text=True)
    return cell


def export(batch_name=None, output=None, *, offline=False, thumbnail_size=240,
           max_skus_per_file=1000, preview_concurrency=8, overwrite=False, dry_run=False):
    if not 100 <= thumbnail_size <= 400 or max_skus_per_file <= 0 or preview_concurrency <= 0:
        raise ValueError('缩略图尺寸须为100～400，分卷商品上限须大于0')
    paths = workflow.batch_paths(batch_name)
    if not paths['marker'].exists():
        raise ValueError('未找到已启动批次，请先创建图片准备计划')
    records = workflow.prepare(batch_name, dry_run=True)
    if not records:
        raise ValueError('批次没有商品记录')
    payloads, manifests = workflow.build_results(records, paths)
    states = {(r['sku'], r['image_name']): r for r in [
        *workflow.load_jsonl(paths['sub_results']), *workflow.load_jsonl(paths['sub_checkpoint'])]}
    local_paths = {}
    for image in manifests:
        if image['origin'] == 'generated' and image.get('url'):
            state = states.get((image['record_id'], image['image_name']), {})
            if state.get('status') == 'success':
                local_paths[image['url']] = state.get('downloaded_path')
    rows = []
    for record, payload in zip(records, payloads):
        refs = [('参考主图', record.get('reference_main_url'))]
        refs.extend((f'参考副图{i}', url) for i, url in enumerate(record.get('reference_secondary_urls', []), 1))
        finals = [('最新使用主图（原样保留）', payload['main_image_url'])]
        finals.extend((f'最新使用副图{i}', url) for i, url in enumerate(payload['secondary_urls'], 1))
        rows.append((record, payload, refs, finals))
    base = Path(output) if output else paths['root'] / '09_review/人工审核预览.xlsx'
    if base.suffix.lower() != '.xlsx':
        raise ValueError('审核输出必须是.xlsx')
    parts = [rows[i:i + max_skus_per_file] for i in range(0, len(rows), max_skus_per_file)]
    outputs = [base] if len(parts) == 1 else [base.with_name(f'{base.stem}_{i:03d}.xlsx') for i in range(1, len(parts) + 1)]
    summary = {'records': len(rows), 'complete': sum(p['complete'] for p in payloads),
               'files': [str(p) for p in outputs], 'embedded_images': 0, 'preview_errors': 0}
    if dry_run:
        return summary
    if not overwrite and any(p.exists() for p in outputs):
        raise ValueError('审核文件已存在；为保护人工意见，请指定新输出路径，或明确使用--overwrite')
    images = PreviewImages(thumbnail_size, offline, preview_concurrency)
    for destination, part in zip(outputs, parts):
        images.preload((url, local_paths.get(url)) for row in part for group in row[2:4] for _, url in group if url)
        reference_count = max(len(row[2]) for row in part)
        final_count = max(len(row[3]) for row in part)
        info_count = len(INFO_HEADERS)
        ref_start, final_start = info_count + 1, info_count + reference_count + 1
        total_columns = info_count + reference_count + final_count
        book = Workbook()
        sheet = book.active
        sheet.title = '人工审核预览'
        for start, end, label, color in (
            (1, info_count, '产品及文案信息 / 人工审核', 'FFF2CC'),
            (ref_start, final_start - 1, '参考图（沃尔玛平台素材）', REFERENCE_COLOR),
            (final_start, total_columns, '最新使用图（原主图＋最终副图）', FINAL_COLOR)):
            sheet.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
            set_cell(sheet, 1, start, label).font = Font(bold=True, size=13)
            for column in range(start, end + 1):
                sheet.cell(1, column).fill = PatternFill('solid', fgColor=color)
        headers = INFO_HEADERS + ['参考主图'] + [f'参考副图{i}' for i in range(1, reference_count)]
        headers += ['最新使用主图（原样保留）'] + [f'最新使用副图{i}' for i in range(1, final_count)]
        for column, header in enumerate(headers, 1):
            cell = set_cell(sheet, 2, column, header)
            cell.font = Font(bold=True)
            color = 'FFF2CC' if column <= info_count else REFERENCE_COLOR if column < final_start else FINAL_COLOR
            cell.fill = PatternFill('solid', fgColor=color)
        sheet.freeze_panes = 'E3'
        sheet.row_dimensions[1].height = 26
        sheet.row_dimensions[2].height = 32
        approval = DataValidation(type='list', formula1='"待审核,通过,退回"')
        approval.errorTitle = '请选择审核结论'
        approval.error = '只能填写待审核、通过或退回'
        approval.showErrorMessage = True
        sheet.add_data_validation(approval)
        for index, (record, payload, refs, finals) in enumerate(part):
            photo_row, link_row = 3 + index * 2, 4 + index * 2
            info = [record['source_sku'], record['result_sku'], record['store'], record['gtin'],
                    record.get('product_type'), record['title'], record['bullets'],
                    payload['complete'], '；'.join(payload['errors']), '待审核', '', '']
            for column, value in enumerate(info, 1):
                sheet.merge_cells(start_row=photo_row, start_column=column, end_row=link_row, end_column=column)
                set_cell(sheet, photo_row, column, value)
            approval.add(sheet.cell(photo_row, INFO_HEADERS.index('审核结论') + 1))
            issues = []
            for start, group, color in ((ref_start, refs, REFERENCE_COLOR), (final_start, finals, FINAL_COLOR)):
                for offset, (label, url) in enumerate(group):
                    column = start + offset
                    for row in (photo_row, link_row):
                        sheet.cell(row, column).fill = PatternFill('solid', fgColor=color)
                    if not url:
                        set_cell(sheet, photo_row, column, '未提供参考图' if start == ref_start else '图片未就绪')
                        continue
                    cell = set_cell(sheet, link_row, column, url)
                    cell.hyperlink = url
                    cell.font = Font(color='0563C1', underline='single', size=9)
                    try:
                        content, width, height = images.get(url, local_paths.get(url))
                        image = ExcelImage(io.BytesIO(content))
                        image.width, image.height = width, height
                        sheet.add_image(image, f'{get_column_letter(column)}{photo_row}')
                        summary['embedded_images'] += 1
                    except Exception as exc:
                        issue = f'{label}：{exc}'
                        issues.append(issue)
                        set_cell(sheet, photo_row, column, '预览加载失败，请打开原图链接核验')
                        summary['preview_errors'] += 1
            set_cell(sheet, photo_row, INFO_HEADERS.index('预览加载问题') + 1, '；'.join(issues))
            sheet.row_dimensions[photo_row].height = thumbnail_size * 0.75 + 8
            sheet.row_dimensions[link_row].height = 44
            print(f"审核预览 [{index + 1}/{len(part)}] {record['result_sku']} | 加载问题={len(issues)}", flush=True)
        widths = [23, 23, 14, 19, 20, 32, 42, 12, 28, 14, 30, 30]
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        for column in range(ref_start, total_columns + 1):
            sheet.column_dimensions[get_column_letter(column)].width = thumbnail_size / 7 + 2
        sheet.sheet_view.zoomScale = 70
        sheet.print_options.horizontalCentered = True
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        guide = book.create_sheet('审核说明')
        instructions = [
            '蓝色区域是输入的沃尔玛参考图；绿色区域是最新使用图，两组主副图顺序分别标注。',
            '每个商品占两行：缩略图在上、原图链接在下；最新使用主图原样保留。',
            '对照来源SKU、结果SKU、店铺和文案，逐项检查商品身份、颜色、形状及副图表达是否正确。',
            '仅成功上传的新图进入最新使用图；未完成商品仍保留，并显示完整状态和错误。',
            '图片预览失败不会用其他图片代替，请查看加载问题并打开原图链接核验。',
            '审核结论默认待审核，人工选择通过或退回并填写意见；预览生成不表示人工审核通过。',
            '该文件仅供人工审核，不是API导入文件；审核确认后才交接最新使用图片URL.xlsx。',
            '不会生成、上传或更新线上商品。再次导出默认拒绝覆盖，避免丢失人工意见。']
        for row, text in enumerate(instructions, 1):
            set_cell(guide, row, 1, text)
            guide.row_dimensions[row].height = 42
        guide.column_dimensions['A'].width = 110
        destination.parent.mkdir(parents=True, exist_ok=True)
        temp = destination.with_suffix('.xlsx.tmp')
        try:
            book.save(temp)
            temp.replace(destination)
        finally:
            book.close()
            if temp.exists():
                temp.unlink()
    return summary


def main():
    parser = argparse.ArgumentParser(description='导出参考图与最新使用图人工审核Excel，不生成或更新图片')
    parser.add_argument('--batch-name', '--batch', dest='batch_name')
    parser.add_argument('--output', type=Path)
    parser.add_argument('--offline', action='store_true', help='只用本地生成图片，不下载远程图片')
    parser.add_argument('--dry-run', action='store_true')
    parser.add_argument('--overwrite', action='store_true', help='明确允许覆盖已存在审核表及人工意见')
    parser.add_argument('--thumbnail-size', type=int, default=240)
    parser.add_argument('--preview-concurrency', type=int, default=8, help='远程预览图并发下载数')
    parser.add_argument('--max-skus-per-file', type=int, default=1000)
    args = parser.parse_args()
    result = export(args.batch_name, args.output, offline=args.offline, dry_run=args.dry_run,
                    overwrite=args.overwrite, thumbnail_size=args.thumbnail_size,
                    max_skus_per_file=args.max_skus_per_file, preview_concurrency=args.preview_concurrency)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
