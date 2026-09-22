"""Persisted replacement plans using the current prompt generation engines."""
from __future__ import annotations
import argparse
import hashlib
import json
import os
import re
import sys
import time
import uuid
from collections import Counter
from dataclasses import asdict, replace
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import unquote, urlsplit
from openpyxl import Workbook, load_workbook

TASK_ROOT = Path(__file__).resolve().parent
PROJECT_ROOT = TASK_ROOT.parents[1]
sys.path.insert(0, str(PROJECT_ROOT / 'src'))
BATCHES_ROOT = TASK_ROOT / 'batches'
TASK_CONFIG = TASK_ROOT / 'config.json'
SCHEMA_VERSION = 3

def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))

def load_task_config():
    return read_json(TASK_CONFIG)

def load_stage_config(stage):
    return read_json(TASK_ROOT / 'stages' / stage / 'config.json')

def load_jsonl(path):
    path = Path(path)
    return [json.loads(s) for s in path.read_text(encoding='utf-8-sig').splitlines() if s.strip()] if path.exists() else []

def save_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + '.tmp')
    tmp.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
    tmp.replace(path)

def save_jsonl(path, rows):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + '.tmp')
    tmp.write_text(''.join(json.dumps(r, ensure_ascii=False) + '\n' for r in rows), encoding='utf-8')
    tmp.replace(path)

def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()

def safe_name(value):
    return re.sub(r'[^\w.-]+', '_', value).strip('._') or digest(value)[:16]

def batch_paths(batch_name=None):
    name = batch_name or Path(load_task_config()['input']['excel_path']).stem
    root = BATCHES_ROOT / safe_name(name)
    paths = {'root': root, 'marker': root / 'replace_batch.json', 'plan': root / '01_inventory/generation_plan.jsonl',
        'inventory': root / '01_inventory/image_inventory.jsonl', 'prompt_tasks': root / '02_prompts/prompt_tasks.jsonl',
        'model_results': root / '02_prompts/model_results.jsonl', 'full_outputs': root / '02_prompts/full_outputs',
        'payload': root / '05_result/image_set.jsonl', 'manifest': root / '05_result/image_manifest.jsonl'}
    for role, dirname in (('main', '04b_generate_main_images'), ('sub', '04_generate_images')):
        folder = root / dirname
        for suffix, filename in (('input', 'image_input.xlsx'), ('results', 'image_generation_results.jsonl'),
            ('checkpoint', 'image_generation_checkpoint.jsonl'), ('excel', 'image_generation_result.xlsx'),
            ('download', 'downloaded_images'), ('raw', 'raw_responses')):
            paths[f'{role}_{suffix}'] = folder / filename
        paths[f'{role}_oss_checkpoint'] = root / f'upload_{role}/oss_upload_checkpoint.jsonl'
        paths[f'{role}_oss_results'] = root / f'upload_{role}/oss_upload_results.jsonl'
    for name, filename in (('latest_used_urls', '最新使用图片URL'),
                           ('new_generated_urls', '本批次新生成图片URL'),
                           ('archive_candidate_urls', '待归档旧副图URL')):
        paths[name] = root / '08_reports' / f'{filename}.jsonl'
    return paths

def task_execution():
    return load_task_config().get('execution', {})

def print_batch_info(batch_name=None):
    print(f"批次目录: {batch_paths(batch_name)['root']}")

def oss_settings():
    from ai_gateway.config.loader import load_local_env
    env = load_local_env(PROJECT_ROOT / 'configs/local.env')
    def get(key, default=''):
        return os.environ.get(key, env.get(key, default)).strip()
    return {'bucket': get('ALIYUN_OSS_BUCKET'), 'endpoint': get('ALIYUN_OSS_ENDPOINT').replace('https://', '').replace('http://', '').rstrip('/'),
            'default_prefix': get('ALIYUN_OSS_DEFAULT_PREFIX', 'images').strip('/')}

def url_parts(url):
    parsed = urlsplit(url)
    if parsed.scheme not in ('https', 'http') or not parsed.hostname or parsed.username or parsed.fragment:
        raise ValueError('无效图片URL')
    key = unquote(parsed.path).lstrip('/')
    if not key or any(p in ('.', '..', '') for p in key.split('/')):
        raise ValueError('无效图片对象路径')
    return parsed.hostname.lower(), key

def classify(url, role):
    _, key = url_parts(url)
    filename = key.rsplit('/', 1)[-1].lower()
    main, sub = 'new_main' in filename, 'new_sub' in filename
    if (main and sub) or (main and role != 'main') or (sub and role != 'sub'):
        raise ValueError('主副图角色或文件名标记冲突')
    return 'new' if main or sub else 'old'

def valid_reference(url):
    host, _ = url_parts(url)
    return any(host == d or host.endswith('.' + d) for d in ('walmartimages.com', 'walmart.com'))

def numbered_columns(row, prefix):
    values = []
    for name, value in row.items():
        match = re.fullmatch(re.escape(prefix) + r'(\d+)', name)
        if match and value:
            values.append((int(match[1]), name, value))
    return [(name, value) for _, name, value in sorted(values)]

def numbered(row, prefix):
    return [value for _, value in numbered_columns(row, prefix)]

def read_input(config):
    path = Path(config['input']['excel_path'])
    if not path.is_absolute():
        path = PROJECT_ROOT / path
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[config['input'].get('sheet_name', 'Sheet1')]
        values = iter(ws.iter_rows(values_only=True))
        header = [str(v).strip() if v is not None else '' for v in next(values, ())]
        for name in ('来源SKU', '结果SKU', '店铺'):
            if name not in header:
                raise ValueError(f'缺少表头: {name}')
        if len([h for h in header if h]) != len(set(h for h in header if h)):
            raise ValueError('表头重复')
        rows = []
        for number, cells in enumerate(values, 2):
            if any(v is not None and str(v).strip() for v in cells):
                rows.append((number, {h: '' if v is None else str(v).strip() for h, v in zip(header, cells) if h}))
        return rows
    finally:
        wb.close()

def analyze(config, settings=None):
    settings = settings if settings is not None else oss_settings()
    data = read_input(config)
    duplicates = Counter((row.get('店铺', ''), row.get('结果SKU', '')) for _, row in data)
    limit = config.get('execution', {}).get('max_records')
    if limit and int(limit) > 0:
        data = data[:int(limit)]
    records = []
    for number, row in data:
        source, result, store = (row.get(k, '') for k in ('来源SKU', '结果SKU', '店铺'))
        r = {'record_id': digest([store, result])[:24], 'row_number': number, 'source_sku': source, 'result_sku': result, 'store': store,
            'gtin': row.get('GTIN') or None, 'product_type': row.get('商品类型') or None, 'title': row.get('标题', ''), 'bullets': row.get('五点', ''),
            'reference_main_url': row.get('参考主图链接', ''), 'reference_secondary_urls': numbered(row, '参考副图链接'),
            'input_fingerprint': digest(row), 'existing': [], 'errors': [], 'oss_directory': None, 'need_main': 0, 'need_sub': 0}
        errors, seen, bases = r['errors'], set(), set()
        if not all((source, result, store)):
            errors.append('来源SKU、结果SKU、店铺不能为空')
        if duplicates[(store, result)] > 1:
            r['record_id'] = digest([store, result, number])[:24]
            errors.append('同店铺结果SKU重复')
        for role, columns in (('main', [('已有主图链接', row.get('已有主图链接', ''))]), ('sub', numbered_columns(row, '已有副图链接'))):
            for column, url in columns:
                input_order = 1 if role == 'main' else int(column.removeprefix('已有副图链接'))
                if not url:
                    continue
                if role == 'main':
                    try:
                        host, key = url_parts(url)
                        match = re.fullmatch(r'([^.]+)\.(oss[^.]*\.aliyuncs\.com)', host)
                        if match and (not settings['bucket'] or match[1] == settings['bucket']):
                            parent = key.rsplit('/', 1)[0]
                            bases.add((match[1], match[2], parent.rsplit('/', 1)[0] if '/' in parent else ''))
                        r['existing'].append({'image_id': digest((host, key))[:24], 'role': 'main',
                            'kind': 'retained', 'url': url, 'oss_key': key, 'origin': 'existing',
                            'input_column': column, 'input_order': input_order,
                            'actual_directory': key.rsplit('/', 1)[0]})
                    except ValueError as exc:
                        errors.append(str(exc))
                    continue
                try:
                    host, key = url_parts(url)
                    match = re.fullmatch(r'([^.]+)\.(oss[^.]*\.aliyuncs\.com)', host)
                    if not match:
                        raise ValueError('已有图片不是OSS公网URL')
                    if settings['bucket'] and match[1] != settings['bucket']:
                        raise ValueError('OSS bucket与上传配置不一致')
                    identity = (match[1], key)
                    kind = classify(url, role)
                    if identity in seen:
                        if any(i['image_id'] == digest(identity)[:24] and i['role'] != role for i in r['existing']):
                            raise ValueError('相同图片主副图角色冲突')
                        continue
                    seen.add(identity)
                    parent = key.rsplit('/', 1)[0]
                    r['existing'].append({'image_id': digest(identity)[:24], 'role': role, 'kind': kind, 'url': url,
                        'oss_key': key, 'origin': 'existing', 'actual_directory': parent,
                        'input_column': column, 'input_order': input_order})
                    base = parent.rsplit('/', 1)[0] if '/' in parent else ''
                    bases.add((match[1], match[2], base))
                except ValueError as exc:
                    errors.append(str(exc))
        if len(bases) == 1:
            bucket, endpoint, base = next(iter(bases))
            if '/' in source or '\\' in source or source in ('.', '..'):
                errors.append('来源SKU不能作为目录名')
            else:
                r['oss_directory'] = {'bucket': bucket, 'endpoint': endpoint, 'prefix': f'{base}/{source}'.strip('/')}
        if not row.get('已有主图链接', ''):
            errors.append('已有主图链接必填，主图原样保留')
        new_sub = [i for i in r['existing'] if i['role'] == 'sub' and i['kind'] == 'new']
        r['need_main'], r['need_sub'] = 0, max(0, 4 - len(new_sub))
        if r['need_main'] or r['need_sub']:
            if not r['oss_directory']:
                errors.append('无法唯一确定OSS原路径基座' if bases else '没有已有OSS链接，无法确定原目录')
            if not r['reference_main_url']:
                errors.append('生成图片需要沃尔玛参考主图')
            for url in [r['reference_main_url'], *r['reference_secondary_urls']]:
                if url:
                    try:
                        if not valid_reference(url):
                            errors.append('参考图必须是沃尔玛平台URL')
                    except ValueError as exc:
                        errors.append(str(exc))
        if r['need_sub'] and not (r['title'] and r['bullets']):
            errors.append('生成副图需要标题和五点')
        records.append(r)
    return records

def prepare(batch_name=None, dry_run=False):
    config, paths = load_task_config(), batch_paths(batch_name)
    provider = config['image_provider']
    if provider not in ('tuzi', 'mxapi'):
        raise ValueError('不支持的图片provider')
    current = analyze(config)
    previous = []
    if paths['marker'].exists():
        marker = read_json(paths['marker'])
        if marker.get('schema_version') != SCHEMA_VERSION or marker.get('provider') != provider:
            raise ValueError('批次布局或平台冲突，请切回或使用新批次')
        previous = load_jsonl(paths['plan'])
        identity = lambda r: (r['row_number'], r['record_id'], r['input_fingerprint'])
        if len(current) < len(previous) or any(identity(old) != identity(new)
                for old, new in zip(previous, current)):
            raise ValueError('input_changed：已有行内容、行号或处理范围变化；原批次只允许末尾追加新行，请恢复原输入或使用新批次')
        current = current[len(previous):]
        if not current:
            return previous
        if previous and current[0]['row_number'] <= previous[-1]['row_number']:
            raise ValueError('input_changed：原批次只允许末尾追加新行')
    elif paths['root'].exists() and any(paths['root'].iterdir()):
        raise ValueError('批次有历史文件但无新布局标记，请使用新批次')
    for r in current:
        r['tasks'] = []
        if r['errors']:
            continue
        for role, count in (('main', r['need_main']), ('sub', r['need_sub'])):
            for index in range(1, count + 1):
                image_id = uuid.uuid4().hex
                now = datetime.now(timezone.utc)
                stamp = now.strftime('%Y%m%dT%H%M%S') + f'{now.microsecond // 1000:03d}Z'
                name = f"new_{role}_{safe_name(r['source_sku'])}_{stamp}_{image_id[:12]}"
                r['tasks'].append({'image_id': image_id, 'role': role, 'ordinal': index, 'image_name': name, 'plan_created_at': now.isoformat(timespec='milliseconds')})
    current = previous + current
    if not dry_run:
        save_jsonl(paths['plan'], current)
        save_jsonl(paths['inventory'], current)
        save_json(paths['marker'], {'schema_version': SCHEMA_VERSION, 'provider': provider})
    return current

def validate_prompt(text, count):
    from ai_gateway.subtasks.walmart_call_prompt_model import inspect_result_text
    from ai_gateway.validators.result_validator import extract_json
    _, _, error = inspect_result_text(text, expected_count=count)
    if error:
        return None, error
    parsed, _ = extract_json(text)
    plan = parsed['image_plan']
    if any(not isinstance(i, dict) for i in plan):
        return None, '提示词计划项不是对象'
    if [i.get('image_number') for i in plan] != list(range(1, count + 1)):
        return None, '提示词编号必须按计划为1..need_sub'
    if any(not isinstance(i.get('ai_image_generation_prompt'), str) or not i['ai_image_generation_prompt'].strip() for i in plan):
        return None, '生成提示词缺失'
    return parsed, None

def uploaded_generated_url(saved, oss_directory, image_name):
    """Accept current JPEG objects and historical PNG objects for one planned image."""
    if not saved or saved.get('status') != 'success' or not saved.get('oss_url') or not saved.get('oss_key'):
        return None
    base = f"{oss_directory['prefix']}/{image_name}"
    key = str(saved['oss_key']).replace('\\', '/')
    if key not in {base + '.jpg', base + '.jpeg', base + '.png'}:
        return None
    host, url_key = url_parts(saved['oss_url'])
    if url_key != key or not host.startswith(oss_directory['bucket'] + '.'):
        return None
    return saved['oss_url']

def prompt_stage(records, paths):
    from ai_gateway.subtasks import walmart_call_prompt_model as shared
    from ai_gateway.clients.openai_chat_client import OpenAIChatClient
    from ai_gateway.config.loader import load_app_config
    tasks = []
    template = (TASK_ROOT / 'prompts/walmart_image_replace_template.txt').read_text(encoding='utf-8-sig')
    for r in records:
        if r['errors'] or not r['need_sub']:
            continue
        prompt = template.replace('{{标题}}', r['title']).replace('{{五点}}', r['bullets']).replace('{{数量}}', str(r['need_sub']))
        tasks.append({'sku': r['record_id'], 'task_id': r['record_id'], 'row_number': r['row_number'],
            'next_task_payload': {'task_id': r['record_id'], 'batch_id': paths['root'].name, 'prompt': prompt,
                'metadata': {'sku': r['record_id'], 'expected_image_plan_count': r['need_sub']},
                'images': [{'type': 'url', 'value': u} for u in [r['reference_main_url'], *r['reference_secondary_urls']]]}})
    save_jsonl(paths['prompt_tasks'], tasks)
    by_id = {r['record_id']: r for r in records}
    existing = {r['sku']: r for r in load_jsonl(paths['model_results'])}
    pending = []
    for task in tasks:
        key = task['sku']
        output = paths['full_outputs'] / f'{key}.json'
        if output.exists():
            _, error = validate_prompt(output.read_text(encoding='utf-8-sig'), by_id[key]['need_sub'])
            if not error:
                existing[key] = {'sku': key, 'task_id': key, 'status': 'success', 'validation_status': 'passed', 'full_output_path': str(output)}
                continue
        recovered = False
        for raw in sorted((paths['full_outputs'] / 'raw').glob(f'{key}__*.json')):
            parsed, error = validate_prompt(raw.read_text(encoding='utf-8-sig'), by_id[key]['need_sub'])
            if not error:
                save_json(output, parsed)
                existing[key] = {'sku': key, 'task_id': key, 'status': 'success', 'validation_status': 'passed', 'full_output_path': str(output)}
                recovered = True
                break
        saved = existing.get(key, {})
        terminal = shared.is_terminal_source_data_error(saved) or shared.is_reference_image_unavailable_error(str(saved.get('error_message') or ''))
        if not recovered and not terminal:
            pending.append(task)
    if pending:
        stage = load_stage_config('generate_prompts')
        model, gateway = stage['execution']['model'], stage['execution']['gateway']['name']
        cfg = shared.WalmartCallPromptModelConfig(name='replace_prompts', input_path=str(paths['prompt_tasks']), output_path=str(paths['model_results']),
            gateways_path=str(PROJECT_ROOT / 'configs/gateways.yaml'), models_path=str(PROJECT_ROOT / 'configs/models.yaml'),
            gateway=gateway, model=model['name'], model_candidates=model.get('candidates', []), max_tokens=model.get('max_tokens', 12000),
            temperature=model.get('temperature'), stream=model.get('stream', False), concurrency=task_execution().get('concurrency', 1),
            full_outputs_dir=str(paths['full_outputs'] / 'raw'), retry_delay_seconds=stage.get('retry', {}).get('retry_delay_seconds', 30),
            request_start_interval_seconds=stage.get('limits', {}).get('request_start_interval_seconds', 1))
        app = load_app_config(cfg.gateways_path, cfg.models_path)
        client = OpenAIChatClient(app.gateways[gateway])
        pool = shared.RuntimeModelPool(app.gateways[gateway], cfg.model, cfg.model_candidates)
        def persist(record):
            value = asdict(record)
            parsed, error = validate_prompt(record.result_text, by_id[record.sku]['need_sub'])
            if record.status == 'success' and not error:
                out = paths['full_outputs'] / f'{record.sku}.json'
                save_json(out, parsed)
                value['full_output_path'] = str(out)
            elif record.status == 'success':
                value.update(status='invalid', validation_status='failed', error_message=error)
            existing[record.sku] = value
            save_jsonl(paths['model_results'], existing.values())
        shared.call_pending_tasks(pending, cfg, client, gateway, pool, on_record=persist)
    save_jsonl(paths['model_results'], existing.values())

def input_rows(records, paths, role):
    rows = []
    uploaded = {(i['sku'], i['image_name']): i for i in [*load_jsonl(paths[f'{role}_oss_results']), *load_jsonl(paths[f'{role}_oss_checkpoint'])]}
    fixed = (PROJECT_ROOT / 'subtasks/walmart_image_prompt/prompts/main_image_optimization_prompt.txt').read_text(encoding='utf-8-sig') if role == 'main' else ''
    for r in records:
        if r['errors']:
            continue
        prompts = {}
        if role == 'sub' and r['need_sub']:
            output = paths['full_outputs'] / f"{r['record_id']}.json"
            if not output.exists():
                continue
            parsed, error = validate_prompt(output.read_text(encoding='utf-8-sig'), r['need_sub'])
            if error:
                continue
            prompts = {p['image_number']: p['ai_image_generation_prompt'] for p in parsed['image_plan']}
        for t in r['tasks']:
            if t['role'] == role:
                saved = uploaded.get((r['record_id'], t['image_name']))
                if uploaded_generated_url(saved, r['oss_directory'], t['image_name']):
                    continue
                rows.append({'sku': r['record_id'], 'reference': [r['reference_main_url'], *r['reference_secondary_urls']],
                    'image_name': t['image_name'], 'prompt': fixed if role == 'main' else prompts[t['ordinal']]})
    return rows

def build_image_input(records, paths, role):
    rows = input_rows(records, paths, role)
    if not rows:
        return []
    count = max(len(row['reference']) for row in rows)
    refs = [f'参考图片{i}' for i in range(1, count + 1)]
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.append(['SKU', *refs, '图片命名', '图片类型', '生成提示词', '下载结果', 'task_id'])
    for row in rows:
        ws.append([row['sku'], *row['reference'], *([''] * (count - len(row['reference']))), row['image_name'],
            'Main Image' if role == 'main' else 'Secondary Image', row['prompt'], '', ''])
    paths[f'{role}_input'].parent.mkdir(parents=True, exist_ok=True)
    wb.save(paths[f'{role}_input'])
    return refs

def image_stage(records, paths, role):
    if role == 'main':
        raise ValueError('本次仅替换副图，禁止生成主图')
    from ai_gateway.subtasks import mxapi_generate_images as shared
    from ai_gateway.clients.image_batch import check_batch_provider
    refs = build_image_input(records, paths, role)
    if not refs:
        return
    provider = load_task_config()['image_provider']
    check_batch_provider(paths['root'], provider, bind=True)
    stage = load_stage_config('generate_main_image' if role == 'main' else 'generate_images')
    cfg_data = dict(stage)
    cfg_data['execution'] = {**stage['execution'], 'gateway': {'name': provider,
        'endpoint_submit': '/v1/videos' if provider == 'tuzi' else '/api/v2/gpt-image-2',
        'endpoint_query': '/v1/videos' if provider == 'tuzi' else '/api/v2/gpt-image/task'}}
    cfg_data['input'] = {'excel_path': str(paths[f'{role}_input']), 'model_results_path': str(paths['model_results']), 'sheet_name': 'Sheet1'}
    cfg_data['output'] = {'excel_path': str(paths[f'{role}_excel']), 'results_path': str(paths[f'{role}_results']),
        'checkpoint_path': str(paths[f'{role}_checkpoint']), 'download_dir': str(paths[f'{role}_download']), 'raw_responses_dir': str(paths[f'{role}_raw'])}
    cfg_data['prompt_mode'] = 'fixed'
    cfg_data['columns'] = {'sku': 'SKU', 'reference_image': refs[0], 'reference_images': refs, 'image_name': '图片命名',
        'image_type': '图片类型', 'prompt': '生成提示词', 'status': '下载结果', 'task_id': 'task_id'}
    cfg = shared.load_config(TASK_ROOT / 'stages/generate_images/config.json', config_data=cfg_data)
    cfg.provider = provider
    cfg.desired_count, cfg.image_type_order, cfg.max_records = None, [], None
    cfg.concurrency = task_execution().get('image_concurrency', 1)
    cfg.max_regenerations_per_image = load_task_config().get('image_generation', {}).get('max_regenerations_per_image', 2)
    cfg.generate_main_images, cfg.generate_sub_images = role == 'main', role == 'sub'
    shared.run(cfg)

class PlanUploadClient:
    """Upload unique planned keys; interrupted uploads may repeat the same key."""
    def __init__(self, client):
        self.client = client
    def public_url(self, key):
        return self.client.public_url(key)
    def upload_file(self, local_path, key, overwrite=False):
        last_error = None
        for attempt in range(1, self.client.config.max_retries + 1):
            try:
                self.client.bucket.put_object_from_file(key, str(local_path))
                return {'success': True, 'size': Path(local_path).stat().st_size}
            except Exception as exc:
                last_error = str(exc)
                if attempt < self.client.config.max_retries:
                    time.sleep(attempt)
        return {'success': False, 'error': last_error}

def upload_stage(records, paths, role):
    if role == 'main':
        raise ValueError('本次仅替换副图，禁止上传主图')
    from ai_gateway.subtasks import oss_upload_images as shared
    from ai_gateway.clients.aliyun_oss_client import AliyunOssClient, load_aliyun_oss_config
    from ai_gateway.subtasks.mxapi_generate_images import is_completed_success
    by_id = {r['record_id']: r for r in records if not r['errors']}
    generated = {(r['sku'], r['image_name']): r for r in [*load_jsonl(paths[f'{role}_results']), *load_jsonl(paths[f'{role}_checkpoint'])]}
    rows = []
    for r in by_id.values():
        for t in r['tasks']:
            image = generated.get((r['record_id'], t['image_name']))
            if t['role'] != role or not image or not is_completed_success(image):
                continue
            rows.append({'sku': r['record_id'], 'image_name': t['image_name'], 'row_number': image['row_number'],
                'local_path': str(Path(image['downloaded_path']).with_suffix('.jpg')),
                'source_path': image['downloaded_path'],
                'oss_key': f"{r['oss_directory']['prefix']}/{t['image_name']}.jpg"})
    checkpoint = shared.CheckpointStore(paths[f'{role}_oss_checkpoint'])
    done = {shared.record_key(r): r for r in checkpoint.rows() if r.get('status') == 'success'}
    pending = []
    for row in rows:
        saved = done.get(shared.row_key(row))
        if saved:
            # Historical PNG successes remain valid after new uploads switch
            # to JPEG; this feature does not migrate existing OSS objects.
            continue
        pending.append(row)
    if pending:
        raw = load_aliyun_oss_config(PROJECT_ROOT)
        if any(r['oss_directory']['bucket'] != raw.bucket for r in by_id.values()):
            raise ValueError('目标OSS bucket与当前配置不同')
        client = PlanUploadClient(AliyunOssClient(replace(raw, default_prefix='')))
        output = (load_task_config().get('oss') or {}).get('image_output') or {}
        cfg = shared.OssUploadConfig(name='replace_upload', project_root=str(PROJECT_ROOT), input_excel_path='', input_sheet_name='Sheet1',
            download_dir='', output_excel_path='', output_results_path=str(paths[f'{role}_oss_results']), checkpoint_path=str(paths[f'{role}_oss_checkpoint']),
            columns={}, oss_prefix='', key_template='{image_name}.{extension}', overwrite=False,
            concurrency=task_execution().get('oss_concurrency', 5), output_format=output.get('format', 'jpeg'),
            jpeg_quality=int(output.get('quality', 95)), jpeg_subsampling=int(output.get('subsampling', 0)),
            jpeg_optimize=bool(output.get('optimize', True)), transparent_policy=output.get('transparent_policy', 'keep_png'),
            delete_source_after_upload=bool(output.get('delete_source_after_upload', True)),
            image_results_path=str(paths[f'{role}_results']))
        shared.process_rows(pending, cfg, client, checkpoint)
    if rows or checkpoint.rows():
        save_jsonl(paths[f'{role}_oss_results'], checkpoint.rows())

def build_results(records, paths):
    uploaded, states = {}, {}
    prompt_results = {row['sku']: row for row in load_jsonl(paths['model_results'])}
    for role in ('main', 'sub'):
        for row in [*load_jsonl(paths[f'{role}_oss_results']), *load_jsonl(paths[f'{role}_oss_checkpoint'])]:
            uploaded[(row['sku'], row['image_name'])] = row
        for row in [*load_jsonl(paths[f'{role}_results']), *load_jsonl(paths[f'{role}_checkpoint'])]:
            states[(row['sku'], row['image_name'])] = row
    payloads, manifests = [], []
    input_rows_by_number = {}
    if any('input_column' not in image for r in records for image in r['existing']):
        input_rows_by_number = dict(read_input(load_task_config()))
    for r in records:
        images = [{**i, 'record_id': r['record_id'], 'source_sku': r['source_sku'], 'result_sku': r['result_sku'], 'store': r['store'],
            'usage_status': 'prepared' if i['kind'] in ('new', 'retained') else 'archive_candidate'} for i in r['existing']]
        original_row = input_rows_by_number.get(r['row_number'], {})
        trustworthy_input = original_row and digest(original_row) == r['input_fingerprint']
        for image in images:
            image.setdefault('image_name', url_parts(image['url'])[1].rsplit('/', 1)[-1])
            image.setdefault('input_column', None)
            image.setdefault('input_order', None)
            image.update(plan_created_at=None, generation_recorded_at=None, upload_recorded_at=None)
            if image['input_column'] is None and trustworthy_input:
                columns = [('已有主图链接', original_row.get('已有主图链接', ''))] if image['role'] == 'main' else numbered_columns(original_row, '已有副图链接')
                for column, value in columns:
                    if value == image['url']:
                        image['input_column'] = column
                        image['input_order'] = 1 if image['role'] == 'main' else int(column.removeprefix('已有副图链接'))
                        break
        task_errors = []
        for t in r.get('tasks', []):
            up = uploaded.get((r['record_id'], t['image_name']))
            state = states.get((r['record_id'], t['image_name']), {})
            url = uploaded_generated_url(up, r['oss_directory'], t['image_name'])
            generation_status = 'success' if url else state.get('status', 'not_submitted')
            expected = up['oss_key'] if url else f"{r['oss_directory']['prefix']}/{t['image_name']}.jpg"
            image = {**t, 'record_id': r['record_id'], 'source_sku': r['source_sku'], 'result_sku': r['result_sku'], 'store': r['store'],
                'origin': 'generated', 'kind': 'new', 'url': url, 'oss_key': expected, 'actual_directory': r['oss_directory']['prefix'],
                'usage_status': 'prepared' if url else 'not_ready', 'generation_status': generation_status,
                'upload_status': (up or {}).get('status', 'not_uploaded'), 'task_id': state.get('task_id'),
                'provider': state.get('provider'), 'attempt': state.get('attempts'),
                'input_column': None, 'input_order': None, 'plan_created_at': t.get('plan_created_at'),
                'generation_recorded_at': state.get('created_at') or None, 'upload_recorded_at': (up or {}).get('created_at') or None, 'error': (up or {}).get('error_message') or state.get('error_message')}
            images.append(image)
            if not url and image['error']:
                task_errors.append(image['error'])
        used = [i for i in images if i['kind'] in ('new', 'retained') and i.get('url')]
        main = [i['url'] for i in used if i['role'] == 'main']
        sub = [i['url'] for i in used if i['role'] == 'sub']
        errors = [*r['errors'], *task_errors]
        if len(main) != 1:
            errors.append('保留主图缺失或不唯一')
        if len(sub) < 4:
            errors.append(f'新副图不足4张：当前{len(sub)}张')
        complete = not r['errors'] and len(main) == 1 and len(sub) >= 4
        prompt_result = prompt_results.get(r['record_id'], {})
        if not complete and r['need_sub'] and prompt_result.get('status') in ('failed', 'failed_permanent', 'invalid'):
            output = paths['full_outputs'] / f"{r['record_id']}.json"
            recovered = False
            if output.exists():
                try:
                    _, validation_error = validate_prompt(output.read_text(encoding='utf-8-sig'), r['need_sub'])
                    recovered = not validation_error
                except (OSError, ValueError, TypeError, KeyError):
                    pass
            if not recovered:
                reason = prompt_result.get('error_message') or prompt_result.get('error_code') or prompt_result['status']
                errors.append(f"generate_prompts: {reason}")
        payloads.append({'schema_version': SCHEMA_VERSION, 'record_id': r['record_id'], 'source_sku': r['source_sku'], 'result_sku': r['result_sku'],
            'store': r['store'], 'gtin': r['gtin'], 'product_type': r.get('product_type'), 'main_image_url': main[0] if len(main) == 1 else None, 'secondary_urls': sub,
            'complete': complete, 'status': 'complete' if complete else ('blocked' if r['errors'] else 'pending'), 'errors': errors,
            'image_set_version': digest([main, sub])})
        for image in images:
            image['image_set_ready'] = complete
        manifests.extend(images)
    return payloads, manifests

def write_reports(payloads, manifests, paths):
    for name, selector in (('latest_used_urls', lambda i: i['kind'] in ('new', 'retained') and i.get('url')),
        ('new_generated_urls', lambda i: i['origin'] == 'generated' and i.get('url')),
        ('archive_candidate_urls', lambda i: i['role'] == 'sub' and i['kind'] == 'old')):
        output = []
        for p in payloads:
            items = [i for i in manifests if i['record_id'] == p['record_id'] and selector(i)]
            main = [i['url'] for i in items if i['role'] == 'main']
            sub = [i['url'] for i in items if i['role'] == 'sub']
            generated = [i for i in manifests if i['record_id'] == p['record_id'] and i['origin'] == 'generated']
            output.append({**p, 'main_image_url': main[0] if main else None, 'secondary_urls': sub,
                'planned_count': len(generated), 'generated_count': sum(i.get('generation_status') == 'success' for i in generated),
                'uploaded_count': sum(bool(i.get('url')) for i in generated),
                'archive_status': 'candidate_only' if name == 'archive_candidate_urls' else None})
        save_jsonl(paths[name], output)
        wb = Workbook()
        ws = wb.active
        ws.title = '图片结果'
        count = max((len(p['secondary_urls']) for p in output), default=0)
        ws.append(['来源SKU', '结果SKU', '店铺', 'GTIN', '商品类型', '主图URL', *[f'副图URL{i}' for i in range(1, count + 1)],
            '图集是否完整', '状态', '失败原因', '计划张数', '生成成功张数', '上传成功张数', '归档状态'])
        for p in output:
            values = [p['source_sku'], p['result_sku'], p['store'], p['gtin'], p.get('product_type'), p['main_image_url'], *p['secondary_urls'],
                *([None] * (count - len(p['secondary_urls']))), p['complete'], p['status'], '；'.join(p['errors']),
                p['planned_count'], p['generated_count'], p['uploaded_count'], p['archive_status']]
            from ai_gateway.subtasks.walmart_call_prompt_model import sanitize_excel_value
            ws.append([sanitize_excel_value(value) for value in values])
        for cells in ws.iter_rows():
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        ws.freeze_panes = 'E2'
        ws.auto_filter.ref = ws.dimensions
        wb.save(paths[name].with_suffix('.xlsx'))

def run_cycle(batch_name=None, dry_run=False, only=None):
    if only in ('generate_main_image', 'upload_main_image'):
        raise ValueError('本次仅替换副图，主图阶段已禁用')
    config = load_task_config()
    switches = config.get('workflow', {})
    records, paths = prepare(batch_name, dry_run), batch_paths(batch_name)
    if dry_run:
        for r in records:
            print(f"{r['source_sku']} -> {r['result_sku']} | 主图补{r['need_main']} 副图补{r['need_sub']} | 目录={r['oss_directory']} | 错误={r['errors']}")
        return {'complete': all(not r['errors'] and not r['need_main'] and not r['need_sub'] for r in records), 'stage_errors': []}
    actions = [('generate_prompts', lambda: prompt_stage(records, paths)),
        ('generate_images', lambda: image_stage(records, paths, 'sub')),
        ('upload_oss', lambda: upload_stage(records, paths, 'sub'))]
    errors = []
    for name, action in actions:
        if (only is None and switches.get(name, True)) or only == name:
            try:
                action()
            except Exception as exc:
                errors.append(f'{name}: {type(exc).__name__}: {exc}')
                print(f'阶段失败，继续输出状态：{errors[-1]}', flush=True)
    payloads, manifests = build_results(records, paths)
    if errors:
        for p in payloads:
            if not p['complete']:
                p['errors'].extend(errors)
    save_jsonl(paths['payload'], payloads)
    save_jsonl(paths['manifest'], manifests)
    write_reports(payloads, manifests, paths)
    print(f"记录{len(records)}；完整{sum(p['complete'] for p in payloads)}；阶段错误{len(errors)}")
    return {'complete': all(p['complete'] for p in payloads) and not errors, 'stage_errors': errors}

def main(only=None):
    parser = argparse.ArgumentParser(description='Walmart旧图替换：准备新图集，不更新线上商品')
    parser.add_argument('--dry-run', action='store_true')
    parser.add_argument('--once', action='store_true')
    parser.add_argument('--batch-name')
    args = parser.parse_args()
    scheduler = load_task_config().get('scheduler', {})
    interval = int(scheduler.get('interval_seconds', 600))
    if interval <= 0:
        raise ValueError('scheduler.interval_seconds必须大于0')
    cycle = 0
    while True:
        cycle += 1
        stats = run_cycle(args.batch_name, args.dry_run, only)
        if only or args.dry_run or args.once or not scheduler.get('enabled', False):
            break
        if scheduler.get('stop_when_complete', True) and stats['complete']:
            break
        if scheduler.get('max_cycles') and cycle >= int(scheduler['max_cycles']):
            break
        try:
            print(f'{interval}秒后续跑；Ctrl+C停止', flush=True)
            time.sleep(interval)
        except KeyboardInterrupt:
            break
