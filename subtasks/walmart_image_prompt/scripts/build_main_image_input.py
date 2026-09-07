"""Build main-image generation input workbook for every SKU.

与副图不同：主图使用固定提示词（不调用 BUZZ），每个 SKU 仅一行：
  - 图片命名 = new_main_{sku}
  - 参考图片链接 = 原主图 URL
  - 生成提示词 = 固定提示词文件内容（prompts/main_image_optimization_prompt.txt）
  - 图片类型 = Main Image

不查 OSS 是否已存在，直接生成；是否续跑由 03 的 checkpoint（skip_success）控制。
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[3] / "src"))

from scripts.build_sub_image_download_input import (
    copy_row_style,
    header_map,
    remove_columns_by_header,
    require_headers,
)

TASK_ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH = TASK_ROOT / "stages" / "build_main_image_input" / "config.json"


def ensure_column(worksheet, headers: dict[str, int], column_name: str, after_header: str | None) -> tuple[dict[str, int], int | None]:
    """确保模板含指定列；不存在则在参考列之后插入。返回 (最新 headers, 列号)。"""
    if column_name in headers:
        return headers, headers[column_name]
    if after_header and after_header in headers:
        anchor = headers[after_header]
    elif headers:
        anchor = max(headers.values())
    else:
        anchor = 1
    worksheet.insert_cols(anchor + 1)
    worksheet.cell(1, anchor + 1, column_name)
    new_headers = header_map(worksheet)
    return new_headers, new_headers.get(column_name)


def read_prompt_file(path: str | Path) -> str:
    p = Path(path)
    if not p.exists():
        raise RuntimeError(f"固定提示词文件不存在: {path}")
    return p.read_text(encoding="utf-8").strip()


def run(config: dict[str, Any]) -> None:
    input_config = config["input"]
    output_config = config["output"]
    columns = config["columns"]

    prompt_file = config.get("prompt_file")
    if not prompt_file:
        raise RuntimeError("build_main_image_input 配置缺少 prompt_file（固定主图提示词路径）")
    fixed_prompt = read_prompt_file(prompt_file)
    if not fixed_prompt:
        raise RuntimeError(f"固定主图提示词为空: {prompt_file}")

    source_wb = load_workbook(input_config["source_excel_path"], read_only=True, data_only=True)
    source_ws = source_wb[input_config["source_sheet_name"]]
    source_headers = header_map(source_ws)
    require_headers(source_headers, [columns["source_sku"], columns["source_main_image"]], "source Excel")

    template_wb = load_workbook(input_config["template_path"])
    template_ws = template_wb[input_config["template_sheet_name"]]
    template_headers = header_map(template_ws)
    require_headers(
        template_headers,
        [
            columns["template_sku"],
            columns["template_reference_image"],
            columns["template_image_name"],
            columns["template_download_result"],
            columns["template_task_id"],
        ],
        "template Excel",
    )
    remove_columns_by_header(template_ws, ["OSS上传结果", "结果确认"])
    template_headers = header_map(template_ws)
    # 确保「生成提示词」「图片类型」列存在
    template_headers, prompt_col = ensure_column(
        template_ws, template_headers, columns.get("template_prompt", "生成提示词"), columns["template_image_name"]
    )
    template_headers, type_col = ensure_column(
        template_ws, template_headers, columns.get("template_image_type", "图片类型"), columns.get("template_prompt", "生成提示词")
    )

    if template_ws.max_row > 1:
        template_ws.delete_rows(2, template_ws.max_row - 1)

    sku_col = source_headers[columns["source_sku"]]
    main_image_col = source_headers[columns["source_main_image"]]
    style_row = 2
    output_row = 2
    generated_skus: list[str] = []
    empty_sku_streak = 0
    stop_after_empty_sku_rows = int(config.get("stop_after_empty_sku_rows", 200))

    main_image_type = config.get("image_type", "Main Image")

    for row_idx in range(2, source_ws.max_row + 1):
        raw_sku = source_ws.cell(row_idx, sku_col).value
        if raw_sku is None:
            empty_sku_streak += 1
            if empty_sku_streak >= stop_after_empty_sku_rows:
                break
            continue
        sku = str(raw_sku).strip()
        if not sku:
            empty_sku_streak += 1
            if empty_sku_streak >= stop_after_empty_sku_rows:
                break
            continue
        empty_sku_streak = 0
        main_image = source_ws.cell(row_idx, main_image_col).value
        main_image = "" if main_image is None else str(main_image).strip()
        if not main_image:
            # 无原主图 URL 无法生成主图，跳过该 SKU
            continue
        generated_skus.append(sku)

        copy_row_style(template_ws, style_row, output_row)
        template_ws.cell(output_row, template_headers[columns["template_sku"]]).value = sku
        template_ws.cell(output_row, template_headers[columns["template_reference_image"]]).value = main_image
        template_ws.cell(output_row, template_headers[columns["template_image_name"]]).value = f"new_main_{sku}"
        template_ws.cell(output_row, template_headers[columns["template_download_result"]]).value = None
        template_ws.cell(output_row, template_headers[columns["template_task_id"]]).value = None
        if prompt_col is not None:
            template_ws.cell(output_row, prompt_col).value = fixed_prompt
        if type_col is not None:
            template_ws.cell(output_row, type_col).value = main_image_type
        output_row += 1

    output_path = Path(output_config["excel_path"])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    template_wb.save(output_path)
    print(f"固定主图提示词: {prompt_file}")
    print(f"generated_skus={len(generated_skus)}")
    print(f"generated_rows={output_row - 2}")
    print(f"output_path={output_path}")


def main() -> None:
    from workflow_common import build_main_image_input_config_for_batch
    config = build_main_image_input_config_for_batch()
    run(config)


if __name__ == "__main__":
    main()
