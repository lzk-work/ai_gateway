#!/usr/bin/env python
"""导出图片复刻审核预览 Excel（手动可选阶段）。"""
from __future__ import annotations

import argparse
import io
import json
from pathlib import Path

import openpyxl
import xlsxwriter
from PIL import Image

from workflow_common import BATCHES_ROOT, batch_name, load_task_config, paths

THUMB_SIZE = 400
JPEG_QUALITY = 75


def read_jsonl(path: Path) -> list[dict]:
    if not path.exists():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def batch_paths(name: str) -> dict[str, Path]:
    root = BATCHES_ROOT / name
    return {
        "root": root,
        "image_results": root / "02_generate_images" / "image_generation_results.jsonl",
        "oss_results": root / "03_upload_oss" / "oss_upload_results.jsonl",
        "review_excel": root / "05_review" / "复刻图片审核预览.xlsx",
    }


def load_source_meta() -> dict[str, dict[str, str]]:
    config = load_task_config()
    source = Path(config["input"]["excel_path"])
    columns = config["input"]["columns"]
    if not source.exists():
        print(f"[提示] 源 Excel 不存在，标题/五点/描述将为空: {source}")
        return {}
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheet_name = config["input"].get("sheet_name")
    sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.active
    iterator = sheet.iter_rows(values_only=True)
    header = next(iterator, ())
    indexes = {str(value).strip(): i for i, value in enumerate(header) if value is not None}

    def cell(row, logical_name: str) -> str:
        index = indexes.get(columns[logical_name])
        return "" if index is None or index >= len(row) or row[index] is None else str(row[index]).strip()

    result: dict[str, dict[str, str]] = {}
    for row in iterator:
        sku = cell(row, "sku")
        if sku:
            result[sku] = {
                "title": cell(row, "title"),
                "features": cell(row, "features"),
                "description": cell(row, "description"),
            }
    workbook.close()
    return result


def image_order(record: dict) -> tuple[int, int, str]:
    name = str(record.get("image_name") or "")
    if name.startswith("new_main_"):
        return (0, 0, name)
    number = record.get("image_number")
    if not isinstance(number, int):
        import re
        match = re.search(r"new_sub(\d+)_", name)
        number = int(match.group(1)) if match else 999
    return (1, number, name)


def thumbnail(path: Path) -> tuple[int, int, io.BytesIO] | None:
    try:
        with Image.open(path) as source:
            source.thumbnail((THUMB_SIZE, THUMB_SIZE))
            output = io.BytesIO()
            source.convert("RGB").save(output, "JPEG", quality=JPEG_QUALITY)
            output.seek(0)
            return source.width, source.height, output
    except Exception as exc:
        print(f"[警告] 缩略图读取失败: {path} -> {exc}")
        return None


def split_skus(skus: list[str], max_skus_per_file: int) -> list[list[str]]:
    if max_skus_per_file <= 0:
        raise ValueError("review.max_skus_per_file 必须大于 0")
    return [skus[index:index + max_skus_per_file] for index in range(0, len(skus), max_skus_per_file)]


def output_paths(base: Path, part_count: int) -> list[Path]:
    if part_count <= 1:
        return [base]
    width = max(3, len(str(part_count)))
    return [base.with_name(f"{base.stem}_{index:0{width}d}{base.suffix}") for index in range(1, part_count + 1)]


def write_review_part(
    output: Path,
    skus: list[str],
    grouped: dict[str, list[dict]],
    oss_urls: dict[str, str],
    metadata: dict[str, dict[str, str]],
    max_images: int,
) -> dict[str, int]:
    image_headers = ["主图", *[f"副图{i}" for i in range(1, max_images)]]
    link_headers = [f"{name}链接" for name in image_headers]
    headers = ["SKU", "标题", "五点", "描述", *image_headers, *link_headers]
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = xlsxwriter.Workbook(output)
    sheet = workbook.add_worksheet("审核预览")
    header_format = workbook.add_format({
        "bold": True, "font_color": "#FFFFFF", "bg_color": "#1F4E78",
        "align": "center", "valign": "vcenter",
    })
    wrap_format = workbook.add_format({"text_wrap": True, "valign": "top"})
    link_format = workbook.add_format({"font_color": "#0563C1", "underline": True})
    missing_format = workbook.add_format({"align": "center", "valign": "vcenter"})
    sheet.write_row(0, 0, headers, header_format)
    sheet.freeze_panes(1, 4)

    image_start = 4
    link_start = image_start + max_images
    embedded = linked = 0
    for row_number, sku in enumerate(skus, start=1):
        meta = metadata.get(sku, {})
        sheet.write(row_number, 0, sku)
        sheet.write(row_number, 1, meta.get("title", ""), wrap_format)
        sheet.write(row_number, 2, meta.get("features", ""), wrap_format)
        sheet.write(row_number, 3, meta.get("description", ""), wrap_format)
        for offset, record in enumerate(grouped[sku]):
            source = Path(str(record.get("downloaded_path") or ""))
            result = thumbnail(source) if source.is_file() else None
            if result:
                _width, _height, data = result
                sheet.embed_image(
                    row_number,
                    image_start + offset,
                    "thumbnail.jpg",
                    {"image_data": data, "description": str(record.get("image_name") or "商品图片")},
                )
                embedded += 1
            else:
                sheet.write(row_number, image_start + offset, "本地图片缺失", missing_format)
            key = f"{sku}::{record.get('image_name')}"
            url = oss_urls.get(key, "")
            if url:
                sheet.write_url(row_number, link_start + offset, url, link_format, url)
                linked += 1
        sheet.set_row(row_number, THUMB_SIZE * 0.75 + 6)

    for column, width in enumerate((22, 35, 45, 45)):
        sheet.set_column(column, column, width)
    sheet.set_column(image_start, link_start - 1, THUMB_SIZE / 7 + 2)
    sheet.set_column(link_start, link_start + max_images - 1, 40)
    sheet.autofilter(0, 0, len(skus), len(headers) - 1)
    workbook.close()
    return {"embedded": embedded, "linked": linked}


def export_review(selected_paths: dict[str, Path], max_skus_per_file: int) -> dict:
    records = [row for row in read_jsonl(selected_paths["image_results"]) if row.get("status") == "success"]
    if not records:
        raise RuntimeError(f"没有可审核的成功图片: {selected_paths['image_results']}")
    grouped: dict[str, list[dict]] = {}
    for record in records:
        sku = str(record.get("sku") or "").strip()
        if sku:
            grouped.setdefault(sku, []).append(record)
    for sku_records in grouped.values():
        sku_records.sort(key=image_order)

    oss_urls: dict[str, str] = {}
    for row in read_jsonl(selected_paths["oss_results"]):
        if row.get("status") not in {"success", "skipped"} or not row.get("oss_url"):
            continue
        oss_urls[f"{row.get('sku')}::{row.get('image_name')}"] = str(row["oss_url"])

    metadata = load_source_meta()
    all_skus = sorted(grouped)
    sku_parts = split_skus(all_skus, max_skus_per_file)
    outputs = output_paths(selected_paths["review_excel"], len(sku_parts))
    base = selected_paths["review_excel"]
    # Remove stale outputs from a previous run with a different number of parts.
    for old in [base, *base.parent.glob(f"{base.stem}_*{base.suffix}")]:
        if old.is_file() and old not in outputs:
            old.unlink()
    embedded = linked = 0
    max_images = max(len(value) for value in grouped.values())
    for output, skus in zip(outputs, sku_parts):
        counts = write_review_part(output, skus, grouped, oss_urls, metadata, max_images)
        embedded += counts["embedded"]
        linked += counts["linked"]
        print(f"分卷完成: {output.name} | SKU={len(skus)} | 图片={counts['embedded']}")
    return {"skus": len(grouped), "embedded": embedded, "linked": linked,
            "files": len(outputs), "outputs": outputs}


def main() -> None:
    parser = argparse.ArgumentParser(description="导出复刻图片审核预览 Excel")
    parser.add_argument("--batch", help="批次目录名；默认使用 config.json 当前输入文件名")
    args = parser.parse_args()
    selected = paths() if not args.batch else batch_paths(args.batch)
    review_config = load_task_config().get("review", {})
    max_skus_per_file = int(review_config.get("max_skus_per_file", 1000))
    result = export_review(selected, max_skus_per_file)
    print("\n=== 复刻图片审核预览导出完成 ===")
    print(
        f"SKU={result['skus']} | 每卷最多={max_skus_per_file} | 文件={result['files']} | "
        f"嵌入图片={result['embedded']} | OSS链接={result['linked']}"
    )
    for output in result["outputs"]:
        print(f"输出: {output}")


if __name__ == "__main__":
    main()
