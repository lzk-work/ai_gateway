from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

OUTPUT_COLUMNS = ["SKU", "处理后主图", *[f"处理后附图{i}" for i in range(1, 7)]]


def read_jsonl(path: Path) -> list[dict]:
    if not path.exists():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def build(image_input: Path, image_results: Path, oss_results: Path, output: Path) -> dict[str, int]:
    source_book = load_workbook(image_input, read_only=True, data_only=True)
    source_sheet = source_book["Sheet1"]
    headers = {str(cell.value).strip(): cell.column for cell in source_sheet[1] if cell.value}
    grouped: dict[str, dict] = {}
    for row_no in range(2, source_sheet.max_row + 1):
        sku = str(source_sheet.cell(row_no, headers["SKU"]).value or "").strip()
        name = str(source_sheet.cell(row_no, headers["图片命名"]).value or "").strip()
        main = str(source_sheet.cell(row_no, headers["产品主图链接"]).value or "").strip()
        if sku and name:
            item = grouped.setdefault(sku, {"source_main": main, "main_name": "", "sub_names": []})
            if name.startswith("new_main_"):
                item["main_name"] = name
            else:
                item["sub_names"].append(name)
    successful_images = {
        f"{row.get('sku')}::{row.get('image_name')}"
        for row in read_jsonl(image_results)
        if row.get("status") == "success"
    }
    uploaded: dict[str, str] = {}
    for row in read_jsonl(oss_results):
        key = f"{row.get('sku')}::{row.get('image_name')}"
        url = str(row.get("oss_url") or "").strip()
        if row.get("status") in {"success", "skipped"} and key in successful_images and url:
            uploaded[key] = url
    rows: list[list[str]] = []
    complete = 0
    for sku, item in grouped.items():
        main_url = uploaded.get(f"{sku}::{item['main_name']}", "")
        urls = [uploaded[f"{sku}::{name}"] for name in item["sub_names"] if f"{sku}::{name}" in uploaded]
        urls = urls[:6]
        if main_url and len(urls) == len(item["sub_names"]):
            complete += 1
        rows.append([sku, main_url, *urls, *([""] * (6 - len(urls)))])
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(OUTPUT_COLUMNS)
    for row in rows:
        sheet.append(row)
    workbook.save(output)
    return {"sku_count": len(rows), "complete_skus": complete, "uploaded_images": len(uploaded)}
