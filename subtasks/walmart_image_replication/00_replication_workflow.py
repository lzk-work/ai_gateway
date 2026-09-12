from __future__ import annotations

import argparse
import json
import runpy
import sys
import time
from pathlib import Path

from workflow_common import TASK_ROOT, load_task_config, paths


def run_step(filename: str, dry_run: bool) -> None:
    old_argv = sys.argv[:]
    sys.argv = [filename] + (["--dry-run"] if dry_run else [])
    try:
        runpy.run_path(str(TASK_ROOT / filename), run_name="__main__")
    finally:
        sys.argv = old_argv


def read_jsonl(path: Path) -> list[dict]:
    if not path.exists():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def complete() -> bool:
    p = paths()
    if not p["image_input"].exists():
        return False
    from openpyxl import load_workbook
    workbook = load_workbook(p["image_input"], read_only=True, data_only=True)
    total = max(workbook["Sheet1"].max_row - 1, 0)
    generated = {f"{r.get('sku')}::{r.get('image_name')}" for r in read_jsonl(p["image_results"]) if r.get("status") == "success"}
    uploaded = {f"{r.get('sku')}::{r.get('image_name')}" for r in read_jsonl(p["oss_results"]) if r.get("status") in {"success", "skipped"}}
    return total > 0 and len(generated) >= total and len(uploaded & generated) >= total


def main() -> None:
    parser = argparse.ArgumentParser(description="Walmart 图片复刻一键流程")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--once", action="store_true")
    args = parser.parse_args()
    cfg = load_task_config()
    stages = [
        ("build_input", "01_build_replication_tasks.py"),
        ("generate_images", "02_generate_replication_images.py"),
        ("upload_oss", "03_upload_replication_images.py"),
        ("build_final_result", "04_build_replication_result.py"),
    ]
    scheduler = cfg.get("scheduler", {})
    interval = int(scheduler.get("interval_seconds", 10800))
    if interval <= 0:
        raise ValueError("scheduler.interval_seconds 必须大于 0")
    max_cycles = scheduler.get("max_cycles")
    max_cycles = int(max_cycles) if max_cycles is not None else None
    cycle = 0
    while True:
        cycle += 1
        print(f"\n=== 图片复刻流程第 {cycle} 轮 ===")
        errors = []
        for switch, filename in stages:
            if not cfg.get("workflow", {}).get(switch, True):
                print(f"跳过 {filename}")
                continue
            try:
                run_step(filename, args.dry_run)
            except KeyboardInterrupt:
                raise
            except Exception as exc:
                errors.append(f"{filename}: {exc}")
                print(f"[本轮阶段失败，继续后续阶段] {filename}: {exc}", flush=True)
        if args.dry_run or args.once or not scheduler.get("enabled", False):
            break
        if not errors and scheduler.get("stop_when_complete", True) and complete():
            print("全部复刻图片已生成并上传，停止续跑。")
            break
        if max_cycles is not None and cycle >= int(max_cycles):
            print(f"已达到 scheduler.max_cycles={max_cycles}，停止周期续跑。")
            break
        print(f"本轮结束，{interval} 秒后重新从构建任务开始；按 Ctrl+C 可安全停止。", flush=True)
        try:
            time.sleep(interval)
        except KeyboardInterrupt:
            print("\n收到中断，已在两轮之间安全停止。")
            break
    print("\n=== 图片复刻总流程结束 ===")


if __name__ == "__main__":
    main()
