from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


OUTPUT_COLUMNS = [
    "SKU", "产品主图链接", "参考副图链接", "图片命名", "图片类型",
    "生成提示词", "下载结果", "task_id",
]


def render_prompt(template: str, title: str, features: str, description: str) -> str:
    return (template.replace("{{标题}}", title)
            .replace("{{五点}}", features)
            .replace("{{描述}}", description))


def preview(config: dict[str, Any], sample_limit: int = 5) -> dict[str, Any]:
    """Read and validate source data without creating batch artifacts."""
    source = Path(config["input"]["excel_path"])
    if not source.exists():
        raise RuntimeError(f"复刻输入 Excel 不存在: {source}")
    sheet_name = config["input"].get("sheet_name", "Sheet1")
    columns = config["input"]["columns"]
    max_sub_images = max(int(config.get("image_generation", {}).get("max_sub_images", 5)), 0)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise RuntimeError(f"复刻输入工作表不存在: {sheet_name}")
        sheet = workbook[sheet_name]
        headers = {str(c.value).strip(): c.column for c in sheet[1] if c.value is not None and str(c.value).strip()}
        required = [columns["sku"], columns["title"], columns["features"], columns["description"],
                    columns["main_image"], *columns["sub_images"]]
        missing = [name for name in required if name not in headers]
        if missing:
            raise RuntimeError("复刻输入缺少列: " + ", ".join(missing))
        sku_count = main_rows = sub_rows = missing_main = 0
        samples: list[dict[str, Any]] = []
        for row_no in range(2, sheet.max_row + 1):
            def value(name: str) -> str:
                raw = sheet.cell(row_no, headers[name]).value
                return "" if raw is None else str(raw).strip()
            sku = value(columns["sku"])
            if not sku:
                continue
            sku_count += 1
            main = value(columns["main_image"])
            if not main:
                missing_main += 1
                continue
            main_rows += 1
            if len(samples) < sample_limit:
                samples.append({"sku": sku, "image_name": f"new_main_{sku}", "reference_count": 1})
            selected_subs = 0
            for index, name in enumerate(columns["sub_images"], start=1):
                if not value(name):
                    continue
                if selected_subs >= max_sub_images:
                    break
                selected_subs += 1
                sub_rows += 1
                if len(samples) < sample_limit:
                    samples.append({"sku": sku, "image_name": f"new_sub{index}_{sku}", "reference_count": 2})
        return {"source_skus": sku_count, "main_rows": main_rows, "sub_rows": sub_rows,
                "generated_rows": main_rows + sub_rows, "missing_main_skus": missing_main, "samples": samples}
    finally:
        workbook.close()


def build(
    config: dict[str, Any],
    output_path: str | Path,
    template_path: str | Path,
    main_prompt_path: str | Path,
) -> dict[str, int]:
    source = Path(config["input"]["excel_path"])
    sheet_name = config["input"].get("sheet_name", "Sheet1")
    columns = config["input"]["columns"]
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    headers = {str(c.value).strip(): c.column for c in sheet[1] if c.value is not None and str(c.value).strip()}
    required = [columns["sku"], columns["title"], columns["features"], columns["description"],
                columns["main_image"], *columns["sub_images"]]
    missing = [name for name in required if name not in headers]
    if missing:
        raise RuntimeError("复刻输入缺少列: " + ", ".join(missing))
    template = Path(template_path).read_text(encoding="utf-8")
    main_prompt = Path(main_prompt_path).read_text(encoding="utf-8").strip()
    if not main_prompt:
        raise RuntimeError(f"主图固定提示词为空: {main_prompt_path}")
    output_rows: list[list[str]] = []
    max_sub_images = max(int(config.get("image_generation", {}).get("max_sub_images", 5)), 0)
    sku_count = 0
    skipped_main = 0
    for row_no in range(2, sheet.max_row + 1):
        def value(name: str) -> str:
            raw = sheet.cell(row_no, headers[name]).value
            return "" if raw is None else str(raw).strip()
        sku = value(columns["sku"])
        if not sku:
            continue
        sku_count += 1
        main_image = value(columns["main_image"])
        if not main_image:
            skipped_main += 1
            continue
        output_rows.append([
            sku, main_image, "", f"new_main_{sku}", "Main Image",
            main_prompt, "", "",
        ])
        prompt = render_prompt(template, value(columns["title"]), value(columns["features"]), value(columns["description"]))
        selected_subs = 0
        for index, source_column in enumerate(columns["sub_images"], start=1):
            sub_image = value(source_column)
            if not sub_image:
                continue
            if selected_subs >= max_sub_images:
                break
            selected_subs += 1
            output_rows.append([
                sku, main_image, sub_image, f"new_sub{index}_{sku}", f"Replication {index}",
                prompt, "", "",
            ])
    workbook.close()
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    out_book = Workbook()
    out_sheet = out_book.active
    out_sheet.title = "Sheet1"
    out_sheet.append(OUTPUT_COLUMNS)
    for row in output_rows:
        out_sheet.append(row)
    out_book.save(output)
    main_rows = sum(1 for row in output_rows if row[4] == "Main Image")
    return {
        "source_skus": sku_count,
        "generated_rows": len(output_rows),
        "main_rows": main_rows,
        "sub_rows": len(output_rows) - main_rows,
        "missing_main_skus": skipped_main,
    }
