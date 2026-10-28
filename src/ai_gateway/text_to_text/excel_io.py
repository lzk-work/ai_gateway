from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any

from .template import normalize_name, stringify


def read_excel(path: Path, sheet_name: str | None, header_row: int = 1) -> tuple[str, list[str], list[tuple[int, dict[str, Any]]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        header_cells = next(sheet.iter_rows(min_row=header_row, max_row=header_row))
        headers = [stringify(cell.value) for cell in header_cells]
        seen: set[str] = set()
        for header in headers:
            key = normalize_name(header)
            if not key:
                continue
            if key in seen:
                raise ValueError(f"Excel 存在重复表头: {header}")
            seen.add(key)
        rows: list[tuple[int, dict[str, Any]]] = []
        for row_number, cells in enumerate(sheet.iter_rows(min_row=header_row + 1), start=header_row + 1):
            values = [cell.value for cell in cells[: len(headers)]]
            row = {header: value for header, value in zip(headers, values) if header}
            if any(stringify(value) for value in row.values()):
                rows.append((row_number, row))
        return sheet.title, headers, rows
    finally:
        workbook.close()


def snapshot_input(source: Path, snapshots: Path, run_id: str) -> Path:
    snapshots.mkdir(parents=True, exist_ok=True)
    destination = snapshots / f"input_{run_id}{source.suffix.lower()}"
    shutil.copy2(source, destination)
    return destination


def export_results(source: Path, destination: Path, sheet_name: str, records: list[dict[str, Any]], output_fields: list[str]) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    workbook = load_workbook(destination)
    sheet = workbook[sheet_name]
    headers = [stringify(cell.value) for cell in sheet[1]]
    extra = [
        "_record_key", "_input_hash", "_status", "_error_category", "_error_message",
        "_model", "_request_id", "_attempts", "_updated_at", "模型返回文本", *output_fields,
    ]
    columns: dict[str, int] = {}
    next_column = len(headers) + 1
    for header in extra:
        if header in columns:
            continue
        if header in headers:
            column = headers.index(header) + 1
        else:
            column = next_column
            next_column += 1
        columns[header] = column
        cell = sheet.cell(1, column, header)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for record in records:
        row = int(record["source_row"])
        values = {
            "_record_key": record.get("record_key"), "_input_hash": record.get("input_hash"),
            "_status": record.get("status"), "_error_category": record.get("error_category"),
            "_error_message": record.get("error_message"), "_model": record.get("model"),
            "_request_id": record.get("request_id"), "_attempts": record.get("attempts"),
            "_updated_at": record.get("updated_at"), "模型返回文本": record.get("raw_text"),
            **(record.get("parsed") or {}),
        }
        for name, value in values.items():
            if name in columns:
                sheet.cell(row, columns[name], value)
    sheet.freeze_panes = sheet.freeze_panes or "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for name, column in columns.items():
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = 22 if name != "_error_message" else 42
    workbook.save(destination)
    workbook.close()
